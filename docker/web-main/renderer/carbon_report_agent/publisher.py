from __future__ import annotations

from collections import defaultdict
from datetime import date
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips

from carbon_report_agent.audit import append_audit_event
from carbon_report_agent.blueprint import ReportBlueprint, load_blueprint
from carbon_report_agent.boilerplate import (
    build_slots,
    load_boilerplate_library,
    render_entry,
    requirements_met,
)
from carbon_report_agent.calculations import CalculationResult
from carbon_report_agent.factors import FactorLibrary
from carbon_report_agent.image_fetch import fetch_image_bytes
from carbon_report_agent.models import ReportInput
from carbon_report_agent.style_pack import (
    add_spacer_paragraphs,
    apply_heading_style,
    apply_paragraph_style,
    apply_section_page,
    approval_label_suffix,
    cover_spacer_count,
    image_slot_align,
    image_slot_width_inches,
    style_approval_table,
)
from carbon_report_agent.template_distiller import TemplatePackage
from carbon_report_agent.toc import (
    add_heading_bookmark,
    build_toc_entries,
    insert_clickable_toc,
)

_FUEL_LABELS = {
    "natural_gas": "\u5929\u7136\u6c14",
    "gasoline": "\u6c7d\u6cb9",
    "diesel": "\u67f4\u6cb9",
}

_MONTH_LABELS = {
    1: "1\u6708",
    2: "2\u6708",
    3: "3\u6708",
    4: "4\u6708",
    5: "5\u6708",
    6: "6\u6708",
    7: "7\u6708",
    8: "8\u6708",
    9: "9\u6708",
    10: "10\u6708",
    11: "11\u6708",
    12: "12\u6708",
}

_CUSTOM_TABLE_ROLES = {
    "workload_monthly",
    "electricity_internal_detail",
    "electricity_balance_detail",
}


def _fuel_unit_for_type(report: ReportInput, fuel_type: str) -> str:
    for record in report.fuel_activity:
        if record.fuel_type == fuel_type:
            return record.unit
    return ""


def _fuel_label(fuel_type: str) -> str:
    return _FUEL_LABELS.get(fuel_type, fuel_type)


def _embed_picture(
    target,
    image_bytes: bytes,
    *,
    width_inches: float,
) -> None:
    """Embed into a Document (add_picture) or a table Cell (inline run)."""
    stream = BytesIO(image_bytes)
    stream.seek(0)
    width = Inches(width_inches)
    # Document has add_picture; table Cell does not (docx.Document is a factory).
    add_picture = getattr(target, "add_picture", None)
    if callable(add_picture):
        add_picture(stream, width=width)
        return
    paragraph = target.paragraphs[0]
    paragraph.text = ""
    run = paragraph.add_run()
    stream.seek(0)
    run.add_picture(stream, width=width)


def _add_picture(
    document: Document,
    image_bytes: bytes,
    *,
    width_inches: float,
    caption: str | None = None,
    align: str = "center",
) -> None:
    before = len(document.paragraphs)
    _embed_picture(document, image_bytes, width_inches=width_inches)
    # Document.add_picture appends a paragraph; style its alignment.
    if len(document.paragraphs) > before:
        pic_para = document.paragraphs[-1]
        if align == "center":
            pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            pic_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            pic_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if caption:
        p = document.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        apply_paragraph_style(p, "body_para")
        # Captions usually not first-line indented
        p.paragraph_format.first_line_indent = Pt(0)


def _add_table_caption(document: Document, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    ppr = paragraph._p.get_or_add_pPr()
    ind = _get_or_add_child(ppr, "w:ind")
    ind.set(qn("w:firstLine"), "482")
    run = paragraph.add_run(caption)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "\u4eff\u5b8b_GB2312")
    fonts.set(qn("w:cs"), "Times New Roman")
    size = _get_or_add_child(rpr, "w:sz")
    size.set(qn("w:val"), "24")
    size_cs = _get_or_add_child(rpr, "w:szCs")
    size_cs.set(qn("w:val"), "24")
    if rpr.find(qn("w:b")) is None:
        rpr.append(OxmlElement("w:b"))
    if rpr.find(qn("w:bCs")) is None:
        rpr.append(OxmlElement("w:bCs"))


def _resolve_image_bytes(report: ReportInput, role: str) -> bytes | None:
    aliases = {role}
    if role == "office_location":
        aliases.add("site_photo")
    elif role == "site_photo":
        aliases.add("office_location")
    image = next((item for item in report.images if item.role in aliases), None)
    if image is None:
        append_audit_event(
            "image_slot_missing",
            {"role": role, "severity": "warning", "reason": "no matching image"},
        )
        return None

    raw: bytes | None = None
    if image.url and image.url.strip():
        raw = fetch_image_bytes(image.url.strip(), role=role)
    else:
        append_audit_event(
            "image_slot_missing",
            {
                "role": role,
                "severity": "warning",
                "reason": "no url",
            },
        )
        return None
    return raw


def _render_image(
    document: Document,
    report: ReportInput,
    role: str | None,
    caption: str | None,
) -> None:
    if not role:
        return
    png_bytes = _resolve_image_bytes(report, role)
    if png_bytes is None:
        return
    image = next(
        (
            item
            for item in report.images
            if item.role == role
            or (role == "office_location" and item.role == "site_photo")
            or (role == "site_photo" and item.role == "office_location")
        ),
        None,
    )
    width = image_slot_width_inches(role)
    align = image_slot_align(role)
    _add_picture(
        document,
        png_bytes,
        width_inches=width,
        caption=caption or (image.caption if image else None),
        align=align,
    )


def _format_compile_date_zh(compile_date: date) -> str:
    return (
        f"{compile_date.year}\u5e74{compile_date.month}\u6708{compile_date.day}\u65e5"
    )


def _format_report_number(report_number: str) -> str:
    number = (report_number or "").strip()
    if not number:
        return ""
    if "\uff1a" in number:
        return number
    if ":" in number:
        return number
    if "-" in number:
        number = number.rsplit("-", 1)[-1]
    return f"\u62a5\u544a\u7f16\u53f7\uff1a{number}"


def _add_styled_paragraph(
    document: Document,
    text: str,
    token_name: str,
) -> None:
    paragraph = document.add_paragraph(text)
    apply_paragraph_style(paragraph, token_name)


def _add_styled_paragraphs(
    document: Document,
    text: str,
    token_name: str,
) -> None:
    for line in str(text).splitlines():
        if line.strip():
            _add_styled_paragraph(document, line.strip(), token_name)


def _omml_r(text: str) -> str:
    return f"<m:r><m:rPr><m:sty m:val=\"p\"/></m:rPr><m:t>{escape(text)}</m:t></m:r>"


def _omml_sub(base: str, sub: str) -> str:
    return (
        "<m:sSub>"
        "<m:e>"
        f"{_omml_r(base)}"
        "</m:e>"
        "<m:sub>"
        f"{_omml_r(sub)}"
        "</m:sub>"
        "</m:sSub>"
    )


def _omml_eco2(sub: str) -> str:
    co2 = _omml_sub("CO", "2")
    return (
        "<m:sSub>"
        "<m:e>"
        f"{_omml_r('E')}"
        "</m:e>"
        "<m:sub>"
        f"{co2}{_omml_r(sub)}"
        "</m:sub>"
        "</m:sSub>"
    )


def _omml_formula_body(key: str) -> str | None:
    formulas = {
        "method_total_formula": "".join(
            [
                _omml_sub("E", "GHG"),
                _omml_r(" = "),
                _omml_sub("E", "\u71c3\u70e7"),
                _omml_r(" + "),
                _omml_sub("E", "\u8fc7\u7a0b"),
                _omml_r(" + "),
                _omml_eco2("\u51c0\u7535"),
                _omml_r(" + "),
                _omml_eco2("\u51c0\u70ed"),
                _omml_r("\uff081\uff09"),
            ]
        ),
        "method_fuel_formula": "".join(
            [
                _omml_eco2("\u71c3\u70e7"),
                _omml_r(" = "),
                _omml_sub("\u2211", "i"),
                _omml_r("\uff08"),
                _omml_sub("AD", "i"),
                _omml_r(" \u00d7 "),
                _omml_sub("CC", "i"),
                _omml_r(" \u00d7 "),
                _omml_sub("OF", "i"),
                _omml_r(" \u00d7 44/12\uff09\uff082\uff09"),
            ]
        ),
        "method_electricity_formula": "".join(
            [
                _omml_eco2("\u51c0\u7535"),
                _omml_r(" = "),
                _omml_sub("AD", "\u51c0\u7535"),
                _omml_r(" \u00d7 "),
                _omml_sub("EF", "\u7535\u529b"),
                _omml_r("\uff083\uff09"),
            ]
        ),
        "method_heat_formula": "".join(
            [
                _omml_eco2("\u51c0\u70ed"),
                _omml_r(" = "),
                _omml_sub("AD", "\u51c0\u70ed"),
                _omml_r(" \u00d7 "),
                _omml_sub("EF", "\u70ed\u529b"),
                _omml_r("\uff084\uff09"),
            ]
        ),
    }
    return formulas.get(key)


def _add_omml_formula_paragraph(document: Document, key: str) -> None:
    body = _omml_formula_body(key)
    if body is None:
        return
    paragraph = document.add_paragraph()
    apply_paragraph_style(paragraph, "formula_para")
    paragraph._p.append(
        parse_xml(
            "<m:oMathPara "
            "xmlns:m=\"http://schemas.openxmlformats.org/officeDocument/2006/math\" "
            "xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\">"
            f"<m:oMath>{body}</m:oMath>"
            "</m:oMathPara>"
        )
    )


def _omml_formula_label_body(label: str) -> str | None:
    labels = {
        "EGHG": _omml_sub("E", "GHG"),
        "E\u71c3\u70e7": _omml_sub("E", "\u71c3\u70e7"),
        "E\u8fc7\u7a0b": _omml_sub("E", "\u8fc7\u7a0b"),
        "ECO2_\u51c0\u7535": _omml_eco2("\u51c0\u7535"),
        "ECO2_\u51c0\u70ed": _omml_eco2("\u51c0\u70ed"),
        "ECO2_\u71c3\u70e7": _omml_eco2("\u71c3\u70e7"),
        "i": _omml_r("i"),
        "ADi,j": _omml_sub("AD", "i,j"),
        "CCi,j": _omml_sub("CC", "i,j"),
        "OFi,j": _omml_sub("OF", "i,j"),
        "AD\u51c0\u7535": _omml_sub("AD", "\u51c0\u7535"),
        "EF\u7535\u529b": _omml_sub("EF", "\u7535\u529b"),
        "AD\u51c0\u70ed": _omml_sub("AD", "\u51c0\u70ed"),
        "EF\u70ed\u529b": _omml_sub("EF", "\u70ed\u529b"),
    }
    return labels.get(label)


def _clear_paragraph_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _set_cell_omml_formula(cell, label: str) -> None:
    body = _omml_formula_label_body(label)
    if body is None:
        return
    paragraph = cell.paragraphs[0]
    _clear_paragraph_content(paragraph)
    paragraph._p.append(
        parse_xml(
            "<m:oMath "
            "xmlns:m=\"http://schemas.openxmlformats.org/officeDocument/2006/math\" "
            "xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\">"
            f"{body}"
            "</m:oMath>"
        )
    )


def _add_styled_empty_paragraphs(
    document: Document,
    count: int,
    token_name: str,
) -> None:
    for _ in range(max(0, int(count))):
        paragraph = document.add_paragraph("")
        apply_paragraph_style(paragraph, token_name)


def _add_front_blank_paragraphs(document: Document, count: int) -> None:
    for _ in range(max(0, int(count))):
        paragraph = document.add_paragraph("")
        paragraph.paragraph_format.first_line_indent = Pt(24)


def _copy_section_page_setup(source, target) -> None:
    target.top_margin = source.top_margin
    target.bottom_margin = source.bottom_margin
    target.left_margin = source.left_margin
    target.right_margin = source.right_margin
    target.header_distance = source.header_distance
    target.footer_distance = source.footer_distance
    target.page_width = source.page_width
    target.page_height = source.page_height
    target.orientation = source.orientation
    target.gutter = source.gutter


def _apply_section_orientation(section, orientation: str | None) -> None:
    if orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Twips(16838)
        section.page_height = Twips(11906)
        section.top_margin = Twips(1803)
        section.bottom_margin = Twips(1803)
        section.left_margin = Twips(1440)
        section.right_margin = Twips(1440)
        section.header_distance = Twips(851)
        section.footer_distance = Twips(992)
    elif orientation == "portrait":
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Twips(11906)
        section.page_height = Twips(16838)
        section.top_margin = Twips(1440)
        section.bottom_margin = Twips(1440)
        section.left_margin = Twips(1803)
        section.right_margin = Twips(1803)
        section.header_distance = Twips(851)
        section.footer_distance = Twips(992)


def _get_or_add_sect_pr_child(section, tag: str):
    sect_pr = section._sectPr
    child = sect_pr.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        sect_pr.append(child)
    return child


def _get_or_add_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _report_table_by_role(report: ReportInput, role: str | None):
    if not role:
        return None
    return next((table for table in report.report_tables if table.role == role), None)


def _has_report_table_role(report: ReportInput, role: str) -> bool:
    payload = _report_table_by_role(report, role)
    return payload is not None and bool(payload.rows)


def _coerce_cell_text(value: str | int | float | None) -> str:
    if value is None:
        return ""
    return str(value)


def _source_names(values) -> str:
    names = []
    for value in values:
        source = getattr(value, "source_name", "") or ""
        if source and source not in names:
            names.append(source)
    return "\u3001".join(names)


def _set_table_layout_fixed(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)
    tbl_width = _get_or_add_child(tbl_pr, "w:tblW")
    tbl_width.set(qn("w:w"), "5000")
    tbl_width.set(qn("w:type"), "pct")
    tbl_jc = _get_or_add_child(tbl_pr, "w:jc")
    tbl_jc.set(qn("w:val"), "center")
    tbl_layout = _get_or_add_child(tbl_pr, "w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    borders = _get_or_add_child(tbl_pr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = _get_or_add_child(borders, f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def _set_table_grid_widths(table) -> None:
    cols = len(table.columns)
    if cols <= 0:
        return
    width = str(max(720, int(9360 / cols)))
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for _ in range(cols):
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), width)
        grid.append(col)
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = _get_or_add_child(tc_pr, "w:tcW")
            tc_w.set(qn("w:w"), width)
            tc_w.set(qn("w:type"), "dxa")


def _style_table_cell_text(cell, *, bold: bool, size_half_points: str = "18") -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if not paragraph.runs:
            paragraph.add_run("")
        for run in paragraph.runs:
            rpr = run._r.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.insert(0, fonts)
            fonts.set(qn("w:ascii"), "Times New Roman")
            fonts.set(qn("w:hAnsi"), "Times New Roman")
            fonts.set(qn("w:eastAsia"), "\u5b8b\u4f53")
            size = _get_or_add_child(rpr, "w:sz")
            size.set(qn("w:val"), size_half_points)
            size_cs = _get_or_add_child(rpr, "w:szCs")
            size_cs.set(qn("w:val"), size_half_points)
            if bold and rpr.find(qn("w:b")) is None:
                rpr.append(OxmlElement("w:b"))
            if bold and rpr.find(qn("w:bCs")) is None:
                rpr.append(OxmlElement("w:bCs"))


def _mark_table_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _style_formal_table(table, *, header_rows: int = 1) -> None:
    _set_table_layout_fixed(table)
    _set_table_grid_widths(table)
    for row_index, row in enumerate(table.rows):
        if row_index < header_rows:
            _mark_table_header_row(row)
        for cell in row.cells:
            _style_table_cell_text(cell, bold=row_index < header_rows)


def _set_table_width_and_layout(
    table,
    *,
    width: str,
    width_type: str,
    layout_type: str,
    jc: str | None = "center",
) -> None:
    if jc == "center":
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
    elif jc == "right":
        table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)
    tbl_style = _get_or_add_child(tbl_pr, "w:tblStyle")
    tbl_style.set(qn("w:val"), "27")
    tbl_width = _get_or_add_child(tbl_pr, "w:tblW")
    tbl_width.set(qn("w:w"), width)
    tbl_width.set(qn("w:type"), width_type)
    existing_jc = tbl_pr.find(qn("w:jc"))
    if jc is None:
        if existing_jc is not None:
            tbl_pr.remove(existing_jc)
    else:
        tbl_jc = _get_or_add_child(tbl_pr, "w:jc")
        tbl_jc.set(qn("w:val"), jc)
    tbl_layout = _get_or_add_child(tbl_pr, "w:tblLayout")
    tbl_layout.set(qn("w:type"), layout_type)
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is not None:
        tbl_pr.remove(borders)
    margins = _get_or_add_child(tbl_pr, "w:tblCellMar")
    for side, value in {"top": "0", "left": "108", "bottom": "0", "right": "108"}.items():
        margin = _get_or_add_child(margins, f"w:{side}")
        margin.set(qn("w:w"), value)
        margin.set(qn("w:type"), "dxa")


def _set_table_grid(table, widths: list[str]) -> None:
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), width)
        grid.append(col)


def _set_table_cell_margins(
    table,
    *,
    top: str = "0",
    left: str = "0",
    bottom: str = "0",
    right: str = "0",
) -> None:
    margins = _get_or_add_child(table._tbl.tblPr, "w:tblCellMar")
    for side, value in {
        "top": top,
        "left": left,
        "bottom": bottom,
        "right": right,
    }.items():
        margin = _get_or_add_child(margins, f"w:{side}")
        margin.set(qn("w:w"), value)
        margin.set(qn("w:type"), "dxa")


def _set_table_borders_single(table, *, color: str = "auto") -> None:
    borders = _get_or_add_child(table._tbl.tblPr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = _get_or_add_child(borders, f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:color"), color)
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")


def _set_row_height_at_least(row, value: str = "397") -> None:
    tr_pr = row._tr.get_or_add_trPr()
    height = _get_or_add_child(tr_pr, "w:trHeight")
    height.set(qn("w:val"), value)
    height.set(qn("w:hRule"), "atLeast")


def _set_cell_width(cell, width: str, width_type: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = _get_or_add_child(tc_pr, "w:tcW")
    tc_w.set(qn("w:w"), width)
    tc_w.set(qn("w:type"), width_type)


def _set_cell_vertical_merge(cell, value: str | None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:vMerge"))
    if existing is not None:
        tc_pr.remove(existing)
    if value is None:
        return
    v_merge = OxmlElement("w:vMerge")
    if value != "continue":
        v_merge.set(qn("w:val"), value)
    tc_pr.append(v_merge)


def _merge_row_cells(
    table,
    *,
    row_index: int,
    start_col: int,
    end_col: int,
    text: str,
    merged_width: str,
    merged_width_type: str = "pct",
    bold: bool = True,
) -> None:
    row = table.rows[row_index]
    merged = row.cells[start_col].merge(row.cells[end_col])
    merged.text = text
    _set_cell_width(merged, merged_width, merged_width_type)
    _set_template_cell_text(
        merged,
        bold=bold,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=True,
    )


def _set_column_vertical_merge(table, col_index: int, start_row: int, end_row: int) -> None:
    for row_index in range(start_row, end_row + 1):
        _set_cell_vertical_merge(
            table.rows[row_index].cells[col_index],
            "restart" if row_index == start_row else "continue",
        )


def _set_border_value(parent, edge: str, value: str) -> None:
    border = _get_or_add_child(parent, f"w:{edge}")
    border.set(qn("w:val"), value)
    if value == "single":
        border.set(qn("w:color"), "auto")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
    else:
        for attr in ("w:color", "w:sz", "w:space"):
            if qn(attr) in border.attrib:
                del border.attrib[qn(attr)]


def _set_cell_borders(
    cell,
    *,
    top: str,
    left: str,
    bottom: str,
    right: str,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = _get_or_add_child(tc_pr, "w:tcBorders")
    _set_border_value(borders, "top", top)
    _set_border_value(borders, "left", left)
    _set_border_value(borders, "bottom", bottom)
    _set_border_value(borders, "right", right)


def _set_template_cell_text(
    cell,
    *,
    bold: bool,
    align: WD_ALIGN_PARAGRAPH | None,
    line_spacing: bool,
) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    v_align = _get_or_add_child(tc_pr, "w:vAlign")
    v_align.set(qn("w:val"), "center")
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        if line_spacing:
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if not paragraph.runs:
            paragraph.add_run("")
        for run in paragraph.runs:
            rpr = run._r.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.insert(0, fonts)
            fonts.set(qn("w:hint"), "default")
            fonts.set(qn("w:ascii"), "Times New Roman")
            fonts.set(qn("w:hAnsi"), "Times New Roman")
            fonts.set(qn("w:eastAsia"), "\u4eff\u5b8b_GB2312")
            fonts.set(qn("w:cs"), "Times New Roman")
            size = _get_or_add_child(rpr, "w:sz")
            size.set(qn("w:val"), "21")
            size_cs = _get_or_add_child(rpr, "w:szCs")
            size_cs.set(qn("w:val"), "21")
            for tag in ("w:b", "w:bCs"):
                existing = rpr.find(qn(tag))
                if existing is not None:
                    rpr.remove(existing)
                if bold:
                    rpr.append(OxmlElement(tag))


def _style_equipment_template_table(table) -> None:
    _set_table_width_and_layout(
        table,
        width="4997",
        width_type="pct",
        layout_type="fixed",
    )
    _set_table_grid(table, ["751", "1145", "2927", "641", "1950", "1103"])
    cell_widths = ["440", "672", "1718", "376", "1144", "647"]
    for row_index, row in enumerate(table.rows):
        _set_row_height_at_least(row)
        for col_index, cell in enumerate(row.cells):
            _set_cell_width(cell, cell_widths[col_index], "pct")
            _set_cell_borders(
                cell,
                top="single",
                left="single" if col_index == 0 else "nil",
                bottom="nil" if row_index == 0 else "single",
                right="single",
            )
            _set_template_cell_text(
                cell,
                bold=False,
                align=None,
                line_spacing=False,
            )


def _style_workload_template_table(table) -> None:
    _set_table_width_and_layout(
        table,
        width="9071",
        width_type="dxa",
        layout_type="autofit",
    )
    _set_table_grid(table, ["3023", "6048"])
    for row_index, row in enumerate(table.rows):
        _set_row_height_at_least(row)
        for col_index, cell in enumerate(row.cells):
            if row_index == 0:
                _set_cell_width(cell, "1666" if col_index == 0 else "3333", "pct")
                top = bottom = "single"
                left = "single" if col_index == 0 else "nil"
                align = None
            else:
                _set_cell_width(
                    cell,
                    "1666" if col_index == 0 else "5682",
                    "pct" if col_index == 0 else "dxa",
                )
                top = "nil" if col_index == 0 else "single"
                bottom = "single"
                left = "single" if not (
                    row_index == len(table.rows) - 1 and col_index == 1
                ) else "nil"
                align = WD_ALIGN_PARAGRAPH.CENTER
            _set_cell_borders(
                cell,
                top=top,
                left=left,
                bottom=bottom,
                right="single",
            )
            _set_template_cell_text(
                cell,
                bold=row_index == 0,
                align=align,
                line_spacing=row_index > 0,
            )


def _style_template_grid_table(
    table,
    *,
    width: str,
    width_type: str,
    layout_type: str,
    grid: list[str],
    cell_widths: list[str] | None = None,
    cell_width_type: str = "pct",
    jc: str | None = None,
    row_height: str | None = "454",
    row_heights: dict[int, str] | None = None,
    header_rows: int = 1,
    body_bold: bool = False,
    align: WD_ALIGN_PARAGRAPH | None = WD_ALIGN_PARAGRAPH.CENTER,
    header_align: WD_ALIGN_PARAGRAPH | None = WD_ALIGN_PARAGRAPH.CENTER,
    borders: bool = True,
) -> None:
    _set_table_width_and_layout(
        table,
        width=width,
        width_type=width_type,
        layout_type=layout_type,
        jc=jc,
    )
    _set_table_grid(table, grid)
    widths = cell_widths if cell_widths is not None else grid
    for row_index, row in enumerate(table.rows):
        height = (row_heights or {}).get(row_index, row_height)
        if height:
            _set_row_height_at_least(row, height)
        if row_index < header_rows:
            _mark_table_header_row(row)
        for col_index, cell in enumerate(row.cells):
            if col_index < len(widths):
                _set_cell_width(cell, widths[col_index], cell_width_type)
            if borders:
                _set_cell_borders(
                    cell,
                    top="single",
                    left="single",
                    bottom="single",
                    right="single",
                )
            _set_template_cell_text(
                cell,
                bold=body_bold or row_index < header_rows,
                align=header_align if row_index < header_rows else align,
                line_spacing=True,
            )


def _style_formula_template_table(table, role: str) -> None:
    grids = {
        "formula_total": ["1095", "345", "6526"],
        "formula_fuel": ["1050", "360", "6556"],
        "formula_electricity": ["1050", "360", "6556"],
        "formula_heat": ["1050", "360", "6556"],
    }
    grid = grids.get(role, ["1050", "360", "6556"])
    _set_table_width_and_layout(
        table,
        width="7966",
        width_type="dxa",
        layout_type="fixed",
        jc="right",
    )
    _set_table_grid(table, grid)
    for row in table.rows:
        for col_index, cell in enumerate(row.cells):
            _set_cell_width(cell, grid[col_index], "dxa")
            _set_template_cell_text(
                cell,
                bold=False,
                align=None,
                line_spacing=False,
            )
            rpr = cell.paragraphs[0].runs[0]._r.get_or_add_rPr()
            for tag in ("w:sz", "w:szCs"):
                node = rpr.find(qn(tag))
                if node is not None:
                    rpr.remove(node)


def _style_emission_sources_template_table(table) -> None:
    if len(table.columns) == 4:
        _style_template_grid_table(
            table,
            width="5000",
            width_type="pct",
            layout_type="autofit",
            grid=["1942", "1299", "2564", "2717"],
            cell_widths=["1139", "762", "1504", "1593"],
            jc=None,
            row_height="660",
        )
    else:
        _style_template_grid_table(
            table,
            width="5000",
            width_type="pct",
            layout_type="autofit",
            grid=["2408", "5925"],
            cell_widths=["1445", "3554"],
            jc=None,
            row_height="454",
        )


def _style_activity_facility_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="4998",
        width_type="pct",
        layout_type="autofit",
        grid=["929", "3882", "3708"],
        cell_widths=["545", "2278", "2176"],
        jc="center",
        row_height="397",
    )


def _style_fuel_gas_monthly_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="4998",
        width_type="pct",
        layout_type="autofit",
        grid=["2408", "5925"],
        cell_widths=["1445", "3554"],
        jc="center",
        row_height="397",
    )


def _style_fuel_liquid_monthly_template_table(table) -> None:
    cols = len(table.columns)
    if cols == 3:
        grid = ["2411", "2950", "2972"]
        cell_widths = ["1447", "1770", "1783"]
    else:
        grid = ["1447"] + ["1770"] * max(0, cols - 1)
        cell_widths = grid
    _style_template_grid_table(
        table,
        width="4998",
        width_type="pct",
        layout_type="autofit",
        grid=grid,
        cell_widths=cell_widths,
        jc=None,
        row_height="454",
    )


def _style_activity_heat_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="9071",
        width_type="dxa",
        layout_type="fixed",
        grid=["2409", "2595", "1423", "2644"],
        cell_widths=["1328", "1430", "784", "1455"],
        jc="center",
        row_height="397",
        body_bold=True,
        align=None,
        header_align=None,
    )


def _style_emission_factors_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="8522",
        width_type="dxa",
        layout_type="fixed",
        grid=["929", "2787", "4806"],
        cell_widths=["929", "2787", "4806"],
        cell_width_type="dxa",
        jc=None,
        row_height="454",
    )
    _set_table_borders_single(table)


def _style_calc_fuel_detail_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="5000",
        width_type="pct",
        layout_type="autofit",
        grid=["1263", "1069", "1311", "1311", "1234", "2334"],
        cell_widths=["741", "627", "769", "769", "723", "1369"],
        jc=None,
        row_height="454",
        header_rows=3,
        align=None,
        header_align=None,
    )


def _style_calc_power_heat_detail_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="5000",
        width_type="pct",
        layout_type="autofit",
        grid=["1507", "1565", "1926", "1964", "1560"],
        cell_widths=["884", "918", "1130", "1151", "915"],
        jc=None,
        row_height="454",
        header_rows=0,
        align=None,
        header_align=None,
    )


def _style_emission_summary_template_table(table) -> None:
    if len(table.columns) == 3:
        grid = ["4871", "1563", "1879"]
        widths = ["2929", "940", "1130"]
    else:
        grid = ["3500", "1200", "1800", "1800"]
        widths = ["2108", "722", "1085", "1085"]
    _style_template_grid_table(
        table,
        width="4992",
        width_type="pct",
        layout_type="autofit",
        grid=grid,
        cell_widths=widths,
        jc=None,
        row_height="454",
        header_rows=0,
    )


def _style_monitoring_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="5199",
        width_type="pct",
        layout_type="fixed",
        grid=["669", "1359", "1621", "1035", "3387", "1769", "1114", "1262", "1262", "1262"],
        cell_widths=["226", "460", "549", "351", "1148", "600", "377", "428", "428", "428"],
        jc="center",
        row_height="397",
    )


def _style_electricity_internal_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="5440",
        width_type="pct",
        layout_type="autofit",
        grid=["735", "2313", "4", "2030", "2310", "1678"],
        cell_widths=["405", "1277", "1277", "1119", "1273", "925"],
        jc="center",
        row_height="397",
        row_heights={1: "192", len(table.rows) - 1: "764"},
        header_rows=2,
        borders=False,
    )
    _set_table_cell_margins(table)
    _set_table_borders_single(table)
    _merge_row_cells(
        table,
        row_index=0,
        start_col=1,
        end_col=5,
        text=table.rows[0].cells[1].text,
        merged_width="4594",
    )


def _style_electricity_balance_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="5440",
        width_type="pct",
        layout_type="autofit",
        grid=["1328", "1722", "2036", "1941", "2043"],
        cell_widths=["732", "2071", "2071", "1070", "1126"],
        jc="center",
        row_height="397",
        row_heights={1: "201", 2: "201", len(table.rows) - 1: "764"},
        header_rows=3,
        borders=False,
    )
    _set_table_cell_margins(table)
    _set_table_borders_single(table)
    _merge_row_cells(
        table,
        row_index=0,
        start_col=1,
        end_col=4,
        text=table.rows[0].cells[1].text,
        merged_width="4267",
    )


def _style_appendix_emission_summary_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="4998",
        width_type="pct",
        layout_type="autofit",
        grid=["2408", "3837", "2082"],
        cell_widths=["3750", "3750", "1249"],
        jc=None,
        row_height="454",
        header_rows=2,
    )


def _style_appendix_fuel_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="4998",
        width_type="pct",
        layout_type="autofit",
        grid=["1507", "1293", "1588", "1503", "1297", "1331"],
        cell_widths=["885", "759", "931", "882", "760", "781"],
        jc=None,
        row_height="454",
        header_rows=3,
        align=None,
        header_align=None,
    )


def _style_appendix_electricity_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="4998",
        width_type="pct",
        layout_type="autofit",
        grid=["777", "3357", "2568", "1817"],
        cell_widths=["456", "1970", "1507", "1066"],
        jc=None,
        row_height="454",
        header_rows=2,
        borders=False,
    )
    _set_table_borders_single(table)


def _style_appendix_heat_template_table(table) -> None:
    _style_template_grid_table(
        table,
        width="4998",
        width_type="pct",
        layout_type="autofit",
        grid=["6014", "2505"],
        cell_widths=["3529", "1470"],
        jc=None,
        row_height=None,
        header_rows=0,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        header_align=WD_ALIGN_PARAGRAPH.LEFT,
        borders=False,
    )
    _set_table_borders_single(table)


def _style_report_table_payload_by_role(table, role: str, header_rows: int) -> None:
    if role == "workload_monthly":
        _style_workload_template_table(table)
    elif role == "electricity_internal_detail":
        _style_electricity_internal_template_table(table)
    elif role == "electricity_balance_detail":
        _style_electricity_balance_template_table(table)
    elif role == "appendix_emission_summary":
        _style_appendix_emission_summary_template_table(table)
    elif role == "appendix_fuel":
        _style_appendix_fuel_template_table(table)
    elif role == "appendix_electricity":
        _style_appendix_electricity_template_table(table)
    elif role == "appendix_heat":
        _style_appendix_heat_template_table(table)
    else:
        _style_formal_table(table, header_rows=header_rows)


def _default_report_table_caption(report: ReportInput, role: str) -> str:
    electricity_caption_numbers = {
        "electricity_internal_detail": "1",
        "electricity_balance_detail": "2",
    }
    number = electricity_caption_numbers.get(role)
    if number is None:
        return ""
    short_name = report.entity_profile.short_name or "\u4f01\u4e1a"
    return (
        f"\u88683.2- {number}  "
        f"{report.metadata.report_year}\u5e74{short_name}"
        "\u51c0\u8d2d\u5165\u7535\u529b\u7edf\u8ba1\u8868\uff08MWh\uff09"
    )


def _render_report_table_payload(
    document: Document,
    report: ReportInput,
    role: str,
) -> bool:
    payload = _report_table_by_role(report, role)
    if payload is None or not payload.rows:
        return False
    cols = max((len(row) for row in payload.rows), default=0)
    if cols <= 0:
        return False
    caption = payload.caption.strip() or _default_report_table_caption(report, role)
    if caption:
        _add_table_caption(document, caption)
    table = document.add_table(rows=0, cols=cols)
    for row_values in payload.rows:
        cells = table.add_row().cells
        for index in range(cols):
            value = row_values[index] if index < len(row_values) else ""
            paragraph = cells[index].paragraphs[0]
            paragraph.text = ""
            paragraph.add_run(_coerce_cell_text(value))
    header_rows = 3 if role in {"appendix_fuel", "electricity_balance_detail"} else 2 if role in {
        "electricity_internal_detail",
        "appendix_emission_summary",
        "appendix_electricity",
    } else 1
    _style_report_table_payload_by_role(table, role, header_rows)
    return True


def _append_page_field(paragraph, cached_text: str) -> None:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)

    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE   \\* MERGEFORMAT"
    instr_run._r.append(instr)

    sep_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(separate)

    paragraph.add_run(cached_text)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def _set_section_footer_page_number(section, cached_text: str) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    for paragraph in list(footer.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)
    paragraph = footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _append_page_field(paragraph, cached_text)


def _configure_body_header_style(document: Document) -> None:
    header_style = document.styles["Header"]._element
    ppr = _get_or_add_child(header_style, "w:pPr")
    p_bdr = _get_or_add_child(ppr, "w:pBdr")
    bottom = _get_or_add_child(p_bdr, "w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:color"), "auto")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    tabs = _get_or_add_child(ppr, "w:tabs")
    for child in list(tabs):
        tabs.remove(child)
    center_tab = OxmlElement("w:tab")
    center_tab.set(qn("w:val"), "center")
    center_tab.set(qn("w:pos"), "4153")
    tabs.append(center_tab)
    right_tab = OxmlElement("w:tab")
    right_tab.set(qn("w:val"), "right")
    right_tab.set(qn("w:pos"), "8306")
    tabs.append(right_tab)
    snap_to_grid = _get_or_add_child(ppr, "w:snapToGrid")
    snap_to_grid.set(qn("w:val"), "0")
    jc = _get_or_add_child(ppr, "w:jc")
    jc.set(qn("w:val"), "center")
    rpr = _get_or_add_child(header_style, "w:rPr")
    size = _get_or_add_child(rpr, "w:sz")
    size.set(qn("w:val"), "18")
    size_cs = _get_or_add_child(rpr, "w:szCs")
    size_cs.set(qn("w:val"), "18")


def _set_section_header_text(section, text: str) -> None:
    header_text = text.strip()
    if not header_text:
        return
    header = section.header
    header.is_linked_to_previous = False
    for paragraph in list(header.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)
    paragraph = header.add_paragraph()
    paragraph.style = "Header"
    ppr = paragraph._p.get_or_add_pPr()
    spacing = _get_or_add_child(ppr, "w:spacing")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    paragraph_rpr = _get_or_add_child(ppr, "w:rPr")
    paragraph_fonts = _get_or_add_child(paragraph_rpr, "w:rFonts")
    paragraph_fonts.set(qn("w:ascii"), "\u4eff\u5b8b_GB2312")
    paragraph_fonts.set(qn("w:eastAsia"), "\u4eff\u5b8b_GB2312")
    paragraph_size = _get_or_add_child(paragraph_rpr, "w:sz")
    paragraph_size.set(qn("w:val"), "18")
    paragraph_size_cs = _get_or_add_child(paragraph_rpr, "w:szCs")
    paragraph_size_cs.set(qn("w:val"), "18")
    run = paragraph.add_run(header_text)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "\u4eff\u5b8b_GB2312")
    fonts.set(qn("w:hAnsi"), "\u4eff\u5b8b_GB2312")
    fonts.set(qn("w:eastAsia"), "\u4eff\u5b8b_GB2312")
    size = _get_or_add_child(rpr, "w:sz")
    size.set(qn("w:val"), "18")
    size_cs = _get_or_add_child(rpr, "w:szCs")
    size_cs.set(qn("w:val"), "18")


def _set_section_page_numbering(
    section,
    *,
    fmt: str | None,
    start: int | None,
) -> None:
    pg_num_type = _get_or_add_sect_pr_child(section, "w:pgNumType")
    if start is not None:
        pg_num_type.set(qn("w:start"), str(start))
    if fmt and fmt != "decimal":
        pg_num_type.set(qn("w:fmt"), fmt)
    elif pg_num_type.get(qn("w:fmt")) is not None:
        del pg_num_type.attrib[qn("w:fmt")]


def _add_section_break(
    document: Document,
    *,
    page_number_format: str | None = None,
    page_number_start: int | None = None,
    header_text: str | None = None,
    section_orientation: str | None = None,
) -> None:
    source_section = document.sections[-1]
    section = document.add_section(WD_SECTION.NEW_PAGE)
    _copy_section_page_setup(source_section, section)
    _apply_section_orientation(section, section_orientation)
    if page_number_format or page_number_start is not None:
        _set_section_page_numbering(
            section,
            fmt=page_number_format,
            start=page_number_start,
        )
        cached_text = "I" if page_number_format == "upperRoman" else "1"
        _set_section_footer_page_number(section, cached_text)
    if header_text:
        _set_section_header_text(section, header_text)


def _append_cover_report_number_textbox(paragraph, text: str) -> None:
    """Add the 2024 sample-style floating report-number textbox."""
    safe_text = escape(text)
    pict = parse_xml(
        f"""
        <w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            xmlns:v="urn:schemas-microsoft-com:vml"
            xmlns:o="urn:schemas-microsoft-com:office:office">
          <v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202"
              path="m,l,21600r21600,l21600,xe">
            <v:stroke joinstyle="miter"/>
            <v:path gradientshapeok="t" o:connecttype="rect"/>
          </v:shapetype>
          <v:shape id="_x0000_s1026" type="#_x0000_t202"
              style="position:absolute;left:0pt;margin-left:271.7pt;margin-top:24.55pt;height:70.3pt;width:146.2pt;mso-position-horizontal-relative:margin;mso-position-vertical-relative:margin;mso-wrap-distance-bottom:0pt;mso-wrap-distance-left:9pt;mso-wrap-distance-right:9pt;mso-wrap-distance-top:0pt;z-index:251659264;mso-width-relative:page;mso-height-relative:page;"
              filled="f" stroked="f">
            <v:textbox inset="0,0,0,0">
              <w:txbxContent>
                <w:p>
                  <w:pPr>
                    <w:jc w:val="left"/>
                  </w:pPr>
                  <w:r>
                    <w:rPr>
                      <w:rFonts w:hint="eastAsia" w:ascii="仿宋" w:hAnsi="仿宋" w:eastAsia="仿宋" w:cs="仿宋"/>
                      <w:b/>
                      <w:bCs/>
                      <w:sz w:val="36"/>
                      <w:szCs w:val="32"/>
                    </w:rPr>
                    <w:t>{safe_text}</w:t>
                  </w:r>
                </w:p>
              </w:txbxContent>
            </v:textbox>
          </v:shape>
        </w:pict>
        """
    )
    run = paragraph.add_run()
    run._r.append(pict)


def _set_table_no_borders(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")
        element.set(qn("w:sz"), "0")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "auto")


def _style_cover_chrome_paragraph(paragraph) -> None:
    paragraph.alignment = None
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = 3.0
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        if rpr.find(qn("w:b")) is None:
            rpr.append(OxmlElement("w:b"))
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        fonts.set(qn("w:eastAsia"), "\u5b8b\u4f53")


def _render_cover_chrome(document: Document, report: ReportInput) -> None:
    chrome_paragraph = None
    if any(item.role == "cover_logo" for item in report.images):
        png = _resolve_image_bytes(report, "cover_logo")
        if png is not None:
            before = len(document.paragraphs)
            _add_picture(
                document,
                png,
                width_inches=image_slot_width_inches("cover_logo"),
                align=image_slot_align("cover_logo"),
            )
            if len(document.paragraphs) > before:
                chrome_paragraph = document.paragraphs[-1]
                _style_cover_chrome_paragraph(chrome_paragraph)
    number = _format_report_number(report.metadata.report_number)
    if number:
        # 2024-007 keeps the report-number textbox in its own blank paragraph
        # after the logo paragraph. Keeping it separate avoids Word expanding
        # the logo paragraph height and pushing cover metadata to page 2.
        add_spacer_paragraphs(document, 1)
        number_paragraph = document.add_paragraph("")
        _append_cover_report_number_textbox(number_paragraph, number)
        _style_cover_chrome_paragraph(number_paragraph)
        add_spacer_paragraphs(document, max(0, cover_spacer_count("after_chrome") - 2))
    else:
        add_spacer_paragraphs(document, cover_spacer_count("after_chrome"))


def _render_cover_meta(document: Document, report: ReportInput) -> None:
    meta = report.metadata
    lines = [
        f"\u62a5\u544a\u4e3b\u4f53\uff1a{meta.entity_name}",
        f"\u62a5\u544a\u5e74\u5ea6\uff1a{meta.report_period}",
        f"\u7f16\u5236\u5355\u4f4d\uff1a{meta.compiler}",
        f"\u7f16\u5236\u65e5\u671f\uff1a{_format_compile_date_zh(meta.compile_date)}",
    ]
    for line in lines:
        _add_styled_paragraph(document, line, "cover_meta")
    add_spacer_paragraphs(document, cover_spacer_count("after_meta"))


def _approval_rows(report: ReportInput) -> list[tuple[str, str]]:
    meta = report.metadata
    return [
        ("\u6279\u51c6", meta.approver),
        ("\u5ba1\u5b9a", meta.validator),
        ("\u5ba1\u6838", meta.reviewer),
        ("\u6821\u6838", meta.checker),
        ("\u7f16\u5199", meta.authors),
    ]


def _render_approval(document: Document, report: ReportInput) -> None:
    rows = _approval_rows(report)
    suffix = approval_label_suffix()
    table = document.add_table(rows=len(rows), cols=2)
    for index, (label, name) in enumerate(rows):
        cells = table.rows[index].cells
        cells[0].text = f"{label}{suffix}"
        cells[1].text = name
    style_approval_table(table)
    add_spacer_paragraphs(document, cover_spacer_count("after_approval"))


def _skip_empty_table(role: str, has_rows: bool) -> bool:
    if has_rows:
        return False
    append_audit_event(
        "table_skipped_empty",
        {"role": role, "severity": "warning"},
    )
    return True


def _render_equipment(document: Document, report: ReportInput) -> None:
    if _skip_empty_table("equipment", bool(report.equipment)):
        return
    table = document.add_table(rows=1, cols=6)
    header = table.rows[0].cells
    header[0].text = "\u5e8f\u53f7"
    header[1].text = "\u4f7f\u7528\u80fd\u6e90"
    header[2].text = "\u4e3b\u8981\u8bbe\u5907\u540d\u79f0"
    header[3].text = "\u6570\u91cf"
    header[4].text = "\u89c4\u683c\u578b\u53f7"
    header[5].text = "\u5907\u6ce8"
    for index, item in enumerate(report.equipment, start=1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = item.energy_type
        cells[2].text = item.name
        cells[3].text = str(item.quantity)
        cells[4].text = item.model
        cells[5].text = item.remark
    _style_equipment_template_table(table)


def _render_emission_sources(document: Document, report: ReportInput) -> None:
    rows = []
    if report.organization_boundary and report.organization_boundary.emission_source_rows:
        rows = list(report.organization_boundary.emission_source_rows)
    if rows:
        table = document.add_table(rows=1, cols=4)
        header = table.rows[0].cells
        header[0].text = "\u6392\u653e\u7c7b\u522b"
        header[1].text = "\u6392\u653e\u8bbe\u65bd"
        header[2].text = "\u4f4d\u7f6e"
        header[3].text = "\u6392\u653e\u6e90"
        for row in rows:
            cells = table.add_row().cells
            cells[0].text = row.category
            cells[1].text = row.facility
            cells[2].text = row.location
            cells[3].text = row.source
        _style_emission_sources_template_table(table)
        return
    sources: list[str] = []
    if report.organization_boundary:
        sources = list(report.organization_boundary.emission_sources)
    if _skip_empty_table("emission_sources", bool(sources)):
        return
    table = document.add_table(rows=1, cols=2)
    header = table.rows[0].cells
    header[0].text = "\u5e8f\u53f7"
    header[1].text = "\u6392\u653e\u6e90"
    for index, source in enumerate(sources, start=1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = source
    _style_emission_sources_template_table(table)


def _render_emission_summary(
    document: Document,
    report: ReportInput,
    calculations: CalculationResult,
) -> None:
    _add_table_caption(
        document,
        f"\u88685.3-1  {report.metadata.report_year}\u5e74\u6e29\u5ba4\u6c14\u4f53\u6392\u653e\u6c47\u603b\u8868",
    )
    table = document.add_table(rows=1, cols=3)
    header = table.rows[0].cells
    header[0].text = "\u6e90\u7c7b\u522b"
    header[1].text = "\u5355\u4f4d"
    header[2].text = f"{report.metadata.report_year}\u5e74"
    rows = [
        (
            "\u5316\u77f3\u71c3\u6599\u71c3\u70e7\u6392\u653e\u91cf",
            calculations.fuel_combustion_tco2,
        ),
        (
            "\u51c0\u8d2d\u5165\u7535\u529b\u9690\u542b\u7684\u6392\u653e\u91cf",
            calculations.electricity_emissions_tco2,
        ),
        (
            "\u51c0\u8d2d\u5165\u70ed\u529b\u9690\u542b\u7684\u6392\u653e\u91cf",
            calculations.heat_emissions_tco2,
        ),
        (
            "\u4f01\u4e1a\u6e29\u5ba4\u6c14\u4f53\u6392\u653e\u603b\u91cf",
            calculations.total_emissions_tco2,
        ),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = "tCO2"
        cells[2].text = f"{value:.2f}"
    _style_emission_summary_template_table(table)


def _render_monitoring(document: Document, report: ReportInput) -> None:
    if _skip_empty_table("monitoring", bool(report.monitoring_devices)):
        return
    _add_table_caption(document, "\u88687-1 \u4f01\u4e1a\u76d1\u6d4b\u8bbe\u5907\u6c47\u603b\u8868")
    table = document.add_table(rows=1, cols=10)
    header = table.rows[0].cells
    headers = [
        "\u5e8f\u53f7",
        "\u5668\u5177\u540d\u79f0",
        "\u89c4\u683c\u578b\u53f7",
        "\u7cbe\u786e\u5ea6\u7b49\u7ea7",
        "\u6d4b\u91cf\u8303\u56f4",
        "\u5b89\u88c5\u5730\u70b9",
        "\u89c4\u5b9a\u7684\u6821\u51c6\u9891\u6b21",
        "\u5b9e\u9645\u7684\u6821\u51c6\u9891\u6b21",
        "\u6821\u51c6\u65f6\u95f4",
        "\u662f\u5426\u7b26\u5408\u6307\u5357\u8981\u6c42",
    ]
    for index, text in enumerate(headers):
        header[index].text = text
    for index, device in enumerate(report.monitoring_devices, start=1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = device.name
        cells[2].text = device.model
        cells[3].text = device.accuracy
        cells[4].text = device.measurement_range
        cells[5].text = device.location
        cells[6].text = device.calibration_frequency
        cells[7].text = device.actual_calibration_frequency
        cells[8].text = device.calibration_time
        cells[9].text = device.guide_compliant
    _style_monitoring_template_table(table)


def _render_activity_facility(
    document: Document,
    report: ReportInput,
) -> None:
    by_facility: dict[str, list[str]] = defaultdict(list)
    for record in report.fuel_activity:
        label = f"{_fuel_label(record.fuel_type)} {record.quantity} {_fuel_unit_for_type(report, record.fuel_type)}"
        by_facility[record.facility].append(label)
    if _skip_empty_table("activity_facility", bool(by_facility)):
        return
    table = document.add_table(rows=1, cols=3)
    header = table.rows[0].cells
    header[0].text = "\u5e8f\u53f7"
    header[1].text = "\u6392\u653e\u8bbe\u65bd"
    header[2].text = "\u6d3b\u52a8\u6c34\u5e73\u6570\u636e"
    for index, (facility, parts) in enumerate(sorted(by_facility.items()), start=1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = facility
        cells[2].text = "\uff1b".join(parts)
    _style_activity_facility_template_table(table)


def _render_workload_fallback(document: Document, report: ReportInput) -> None:
    if _skip_empty_table("workload_monthly", bool(report.workload)):
        return
    table = document.add_table(rows=1, cols=4)
    header = table.rows[0].cells
    header[0].text = "\u5e8f\u53f7"
    header[1].text = "\u6307\u6807"
    header[2].text = "\u6570\u91cf"
    header[3].text = "\u5355\u4f4d"
    for index, record in enumerate(report.workload, start=1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = record.name
        cells[2].text = str(record.quantity)
        cells[3].text = record.unit
    _style_formal_table(table)


def _has_month_fuel(report: ReportInput) -> bool:
    return any(r.month is not None for r in report.fuel_activity)


def _render_activity_fuel(
    document: Document,
    report: ReportInput,
    calculations: CalculationResult,
) -> None:
    if _has_month_fuel(report):
        _render_activity_fuel_monthly(document, report)
        return
    liquid_totals = {
        fuel_type: quantity
        for fuel_type, quantity in calculations.activity_fuel_totals.items()
        if fuel_type in {"gasoline", "diesel"}
    }
    if _skip_empty_table("activity_fuel", bool(liquid_totals)):
        return
    table = document.add_table(rows=1, cols=3)
    header = table.rows[0].cells
    header[0].text = "\u71c3\u6599\u7c7b\u578b"
    header[1].text = "\u5e74\u6d3b\u52a8\u91cf"
    header[2].text = "\u5355\u4f4d"
    for fuel_type, quantity in sorted(liquid_totals.items()):
        cells = table.add_row().cells
        cells[0].text = _fuel_label(fuel_type)
        cells[1].text = str(quantity)
        cells[2].text = _fuel_unit_for_type(report, fuel_type)
    _style_template_grid_table(
        table,
        width="4998",
        width_type="pct",
        layout_type="autofit",
        grid=["2408", "2950", "2972"],
        cell_widths=["1447", "1770", "1783"],
        jc=None,
        row_height="454",
    )


def _render_activity_natural_gas(document: Document, report: ReportInput) -> None:
    gas_records = [record for record in report.fuel_activity if record.fuel_type == "natural_gas"]
    if _skip_empty_table("activity_natural_gas", bool(gas_records)):
        return
    gas_by_month: dict[int, float] = defaultdict(float)
    for record in gas_records:
        if record.month is not None:
            gas_by_month[record.month] += record.quantity

    short_name = report.entity_profile.short_name or "\u4f01\u4e1a"
    _add_table_caption(
        document,
        f"\u88683.1- 2  {report.metadata.report_year}\u5e74{short_name}\u5929\u7136\u6c14\u6d88\u8017\u91cf\u7edf\u8ba1\u8868\uff08\u4e07m3\uff09",
    )
    table = document.add_table(rows=1, cols=2)
    header = table.rows[0].cells
    if gas_by_month:
        header[0].text = "\u6708\u4efd"
        header[1].text = f"{report.metadata.report_year}\u5e74\uff08\u4e07Nm\u00b3\uff09"
        total = 0.0
        for month in range(1, 13):
            if month not in gas_by_month:
                continue
            cells = table.add_row().cells
            cells[0].text = _MONTH_LABELS[month]
            cells[1].text = f"{gas_by_month[month]:.4f}"
            total += gas_by_month[month]
        cells = table.add_row().cells
        cells[0].text = "\u5408\u8ba1"
        cells[1].text = f"{total:.4f}"
    else:
        header[0].text = "\u9879\u76ee"
        header[1].text = f"{report.metadata.report_year}\u5e74\uff08\u4e07Nm\u00b3\uff09"
        cells = table.add_row().cells
        cells[0].text = "\u5929\u7136\u6c14"
        cells[1].text = f"{sum(record.quantity for record in gas_records):.4f}"
    _style_fuel_gas_monthly_template_table(table)


def _render_activity_fuel_monthly(document: Document, report: ReportInput) -> None:
    gas_diesel: dict[int, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for record in report.fuel_activity:
        if record.month is None or record.fuel_type not in ("gasoline", "diesel"):
            continue
        key = record.fuel_type
        if record.fuel_type == "diesel":
            key = f"diesel:{record.facility}"
        gas_diesel[record.month][key] += record.quantity

    if gas_diesel:
        table = document.add_table(rows=1, cols=5)
        header = table.rows[0].cells
        header[0].text = "\u6708\u4efd"
        header[1].text = "\u6c7d\u6cb9(t)"
        header[2].text = "\u6c7d\u8f66\u7528\u67f4\u6cb9(t)"
        header[3].text = "\u8f68\u9053\u8f66\u7528\u67f4\u6cb9(t)"
        header[4].text = "\u5408\u8ba1(t)"
        totals = [0.0, 0.0, 0.0, 0.0]
        for month in range(1, 13):
            if month not in gas_diesel:
                continue
            bucket = gas_diesel[month]
            gas = bucket.get("gasoline", 0.0)
            diesel_vehicle = bucket.get("diesel:vehicle", 0.0)
            diesel_rail = bucket.get("diesel:rail-work-vehicle", 0.0)
            # Any other diesel facility rolls into rail column for display
            other_diesel = sum(
                v for k, v in bucket.items() if k.startswith("diesel:") and k not in (
                    "diesel:vehicle",
                    "diesel:rail-work-vehicle",
                )
            )
            diesel_rail += other_diesel
            row_total = gas + diesel_vehicle + diesel_rail
            cells = table.add_row().cells
            cells[0].text = _MONTH_LABELS[month]
            cells[1].text = f"{gas:.4f}"
            cells[2].text = f"{diesel_vehicle:.4f}"
            cells[3].text = f"{diesel_rail:.4f}"
            cells[4].text = f"{row_total:.4f}"
            totals[0] += gas
            totals[1] += diesel_vehicle
            totals[2] += diesel_rail
            totals[3] += row_total
        cells = table.add_row().cells
        cells[0].text = "\u5408\u8ba1"
        cells[1].text = f"{totals[0]:.4f}"
        cells[2].text = f"{totals[1]:.4f}"
        cells[3].text = f"{totals[2]:.4f}"
        cells[4].text = f"{totals[3]:.4f}"
        _style_fuel_liquid_monthly_template_table(table)


def _render_activity_electricity_monthly(
    document: Document,
    report: ReportInput,
) -> None:
    by_month: dict[int, float] = defaultdict(float)
    for record in report.electricity_activity:
        if record.month is None:
            continue
        by_month[record.month] += record.quantity_mwh
    if not by_month:
        total = sum(r.quantity_mwh for r in report.electricity_activity)
        if _skip_empty_table("activity_electricity", total > 0):
            return
        table = document.add_table(rows=1, cols=2)
        header = table.rows[0].cells
        header[0].text = "\u9879\u76ee"
        header[1].text = "MWh"
        cells = table.add_row().cells
        cells[0].text = "\u51c0\u8d2d\u5165\u7535\u529b"
        cells[1].text = f"{total:.3f}"
        _style_fuel_gas_monthly_template_table(table)
        return
    table = document.add_table(rows=1, cols=2)
    header = table.rows[0].cells
    header[0].text = "\u6708\u4efd"
    header[1].text = f"{report.metadata.report_year}\u5e74\uff08MWh\uff09"
    total = 0.0
    for month in range(1, 13):
        if month not in by_month:
            continue
        cells = table.add_row().cells
        cells[0].text = _MONTH_LABELS[month]
        cells[1].text = f"{by_month[month]:.3f}"
        total += by_month[month]
    cells = table.add_row().cells
    cells[0].text = "\u5408\u8ba1"
    cells[1].text = f"{total:.3f}"
    _style_fuel_gas_monthly_template_table(table)


def _render_activity_heat_sites(document: Document, report: ReportInput) -> None:
    if _skip_empty_table("activity_heat", bool(report.heat_activity)):
        return
    _add_table_caption(document, "\u88683.3- 1  \u4f01\u4e1a\u51c0\u8d2d\u5165\u70ed\u529b\u7edf\u8ba1\u8868")
    table = document.add_table(rows=2, cols=4)
    header = table.rows[0].cells
    header[0].text = "\u7ad9\u70b9"
    header[1].text = "\u4f9b\u6696\u9762\u79ef\uff08m2)"
    header[2].text = "\u8017\u70ed\u91cf\u7cfb\u6570"
    header[3].text = "\u70ed\u529b\uff08GJ\uff09"
    units = table.rows[1].cells
    units[0].text = ""
    units[1].text = f"{report.metadata.report_year}\u5e74"
    units[2].text = "GJ/m2"
    units[3].text = f"{report.metadata.report_year}\u5e74"
    for record in report.heat_activity:
        cells = table.add_row().cells
        cells[0].text = record.site
        cells[1].text = (
            f"{record.heating_area_m2:.2f}" if record.heating_area_m2 is not None else ""
        )
        cells[2].text = (
            f"{record.heat_coefficient:.4f}" if record.heat_coefficient is not None else ""
        )
        cells[3].text = f"{record.quantity_gj:.2f}"
    source = table.add_row().cells
    source[0].text = "\u6570\u636e\u6765\u6e90"
    source[1].text = _source_names(report.heat_activity)
    _style_activity_heat_template_table(table)
    _merge_row_cells(
        table,
        row_index=len(table.rows) - 1,
        start_col=1,
        end_col=3,
        text=_source_names(report.heat_activity),
        merged_width="5662",
        merged_width_type="dxa",
    )


def _render_emission_factors(
    document: Document,
    report: ReportInput,
    factors: FactorLibrary,
) -> None:
    del factors
    _add_table_caption(
        document,
        "\u88684.1- 1 \u71c3\u6599\u71c3\u70e7CO2\u6392\u653e\u7684\u6392\u653e\u56e0\u5b50\u5217\u8868",
    )
    table = document.add_table(rows=1, cols=3)
    header = table.rows[0].cells
    header[0].text = "\u5e8f\u53f7"
    header[1].text = "\u6392\u653e\u8bbe\u65bd"
    header[2].text = "\u6392\u653e\u56e0\u5b50"

    by_facility: dict[str, set[str]] = defaultdict(set)
    for record in report.fuel_activity:
        if record.fuel_type in {"natural_gas", "gasoline", "diesel"}:
            by_facility[record.facility].add(record.fuel_type)

    groups: list[tuple[str, list[str]]] = []
    if by_facility.get("vehicle"):
        fuels = [fuel for fuel in ("gasoline", "diesel") if fuel in by_facility["vehicle"]]
        if fuels:
            groups.append(("\u516c\u52a1\u8f66", fuels))
    if by_facility.get("canteen"):
        groups.append(("\u71c3\u6c14\u7076\u5177", ["natural_gas"]))
    rail_fuels = set()
    for facility, fuels in by_facility.items():
        if facility not in {"vehicle", "canteen"}:
            rail_fuels.update(fuel for fuel in fuels if fuel == "diesel")
    if rail_fuels:
        groups.append(("\u5185\u71c3\u673a\u8f66", ["diesel"]))
    if not groups:
        fuel_types = {record.fuel_type for record in report.fuel_activity}
        fallback_fuels = [fuel for fuel in ("gasoline", "diesel") if fuel in fuel_types]
        if fallback_fuels:
            groups.append(("\u516c\u52a1\u8f66", fallback_fuels))
        if "natural_gas" in fuel_types:
            groups.append(("\u71c3\u6c14\u7076\u5177", ["natural_gas"]))

    index = 1
    merge_ranges: list[tuple[int, int]] = []
    for facility_label, fuel_types in groups:
        if not fuel_types:
            continue
        start_row = len(table.rows)
        row_count = 0
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = facility_label
        fuel_type = fuel_types[0]
        cells[2].text = f"{_fuel_label(fuel_type)}\u542b\u78b3\u91cf"
        index += 1
        row_count += 1
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = ""
        cells[2].text = f"{_fuel_label(fuel_type)}\u78b3\u6c27\u5316\u7387"
        index += 1
        row_count += 1
        for fuel_type in fuel_types[1:]:
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = ""
            cells[2].text = f"{_fuel_label(fuel_type)}\u542b\u78b3\u91cf"
            index += 1
            row_count += 1
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = ""
            cells[2].text = f"{_fuel_label(fuel_type)}\u78b3\u6c27\u5316\u7387"
            index += 1
            row_count += 1
        merge_ranges.append((start_row, row_count))
    _style_emission_factors_template_table(table)
    for start_row, row_count in merge_ranges:
        for offset in range(row_count):
            _set_cell_vertical_merge(
                table.rows[start_row + offset].cells[1],
                "restart" if offset == 0 else "continue",
            )


_FUEL_FACTOR_CONSTANTS = {
    "natural_gas": {
        "ncv": "389.31",
        "unit_carbon": "15.30",
        "unit": "\u4e07Nm3",
        "source_prefix": "\u6309\u7167",
        "verb": "\u8ba1\u7b97\u5f97\u5230",
    },
    "gasoline": {
        "ncv": "44.800",
        "unit_carbon": "18.9",
        "unit": "t",
        "source_prefix": "\u6839\u636e",
        "verb": "\u8ba1\u7b97\u5f97\u51fa",
    },
    "diesel": {
        "ncv": "43.330",
        "unit_carbon": "20.20",
        "unit": "t",
        "source_prefix": "\u6839\u636e",
        "verb": "\u8ba1\u7b97\u5f97\u51fa",
    },
}


def _used_factor_fuel_types(report: ReportInput) -> list[str]:
    used = {record.fuel_type for record in report.fuel_activity}
    return [fuel_type for fuel_type in ("natural_gas", "gasoline", "diesel") if fuel_type in used]


def _fuel_factor_heading_text(index: int, fuel_type: str, suffix: str) -> str:
    return f"4.1.{index} {_fuel_label(fuel_type)}\u7684{suffix}"


def _format_percent(rate: float) -> str:
    return f"{rate * 100:.0f}%"


def _add_heading_with_optional_bookmark(
    document: Document,
    text: str,
    level: int,
    bookmark_by_text_level: dict[tuple[int, str], object],
    next_bookmark_id: int,
) -> int:
    paragraph = document.add_heading(text, level=level)
    apply_heading_style(paragraph, level)
    entry = bookmark_by_text_level.get((level, text.strip()))
    if entry is not None:
        add_heading_bookmark(paragraph, entry.bookmark, next_bookmark_id)
        next_bookmark_id += 1
    return next_bookmark_id


def _render_fuel_factor_sections(
    document: Document,
    report: ReportInput,
    factors: FactorLibrary,
    bookmark_by_text_level: dict[tuple[int, str], object],
    next_bookmark_id: int,
) -> int:
    short_name = report.entity_profile.short_name or report.metadata.entity_name or "\u4f01\u4e1a"
    fuel_types = _used_factor_fuel_types(report)
    index = 1
    for fuel_type in fuel_types:
        factor = factors.fuel_factors[fuel_type]
        constants = _FUEL_FACTOR_CONSTANTS[fuel_type]
        label = _fuel_label(fuel_type)
        next_bookmark_id = _add_heading_with_optional_bookmark(
            document,
            _fuel_factor_heading_text(index, fuel_type, "\u542b\u78b3\u91cf"),
            3,
            bookmark_by_text_level,
            next_bookmark_id,
        )
        _add_styled_paragraph(
            document,
            f"{short_name}\u672a\u5bf9{label}\u7684\u542b\u78b3\u91cf\u3001"
            "\u4f4e\u4f4d\u53d1\u70ed\u91cf\u548c\u5355\u4f4d\u70ed\u503c\u542b\u78b3\u91cf\u5f00\u5c55\u5b9e\u6d4b\uff0c"
            f"\u542b\u78b3\u91cf\u91c7\u7528\u300a\u8fd0\u8f93\u4f01\u4e1a\u6838\u7b97\u6307\u5357\u300b\u4e2d{label}\u7684"
            "\u4f4e\u4f4d\u53d1\u70ed\u91cf\u4e0e\u5355\u4f4d\u70ed\u503c\u542b\u78b3\u91cf\u7f3a\u7701\u503c\u76f8\u4e58\u8ba1\u7b97\u5f97\u5230\u3002",
            "body_para",
        )
        _add_styled_paragraph(
            document,
            f"{constants['source_prefix']}\u300a\u8fd0\u8f93\u4f01\u4e1a\u6838\u7b97\u6307\u5357\u300b\uff0c"
            f"{label}\u7684\u4f4e\u4f4d\u53d1\u70ed\u91cf\u548c\u5355\u4f4d\u70ed\u503c\u542b\u78b3\u91cf"
            f"{'\u7f3a\u7701\u503c' if fuel_type != 'natural_gas' else ''}\u5206\u522b\u4e3a"
            f"{constants['ncv']} GJ/{constants['unit']}\u548c{constants['unit_carbon']}\u00d710-3 tC/GJ\uff0c"
            f"{constants['verb']}{label}\u7684\u542b\u78b3\u91cf\u4e3a"
            f"{factor.carbon_content:.4f} tC/{constants['unit']}\u3002",
            "body_para",
        )
        index += 1
    for fuel_type in fuel_types:
        factor = factors.fuel_factors[fuel_type]
        label = _fuel_label(fuel_type)
        next_bookmark_id = _add_heading_with_optional_bookmark(
            document,
            _fuel_factor_heading_text(index, fuel_type, "\u78b3\u6c27\u5316\u7387"),
            3,
            bookmark_by_text_level,
            next_bookmark_id,
        )
        _add_styled_paragraph(
            document,
            f"{short_name}\u672a\u5bf9{label}\u7684\u78b3\u6c27\u5316\u7387\u5f00\u5c55\u5b9e\u6d4b\uff0c"
            f"\u6839\u636e\u300a\u8fd0\u8f93\u4f01\u4e1a\u6838\u7b97\u6307\u5357\u300b\u7684\u89c4\u5b9a\uff0c"
            f"{label}\u7684\u78b3\u6c27\u5316\u7387\u53d6\u7f3a\u7701\u503c{_format_percent(factor.oxidation_rate)}\u3002",
            "body_para",
        )
        index += 1
    return next_bookmark_id


def _is_static_fuel_factor_heading(level: int, text: str) -> bool:
    stripped = text.strip()
    return level == 3 and any(stripped.startswith(f"4.1.{index} ") for index in range(1, 7))


def _chapter4_grid_factor_text(report: ReportInput, factors: FactorLibrary) -> str:
    if not report.electricity_activity:
        return ""
    short_name = report.entity_profile.short_name or report.metadata.entity_name or "\u4f01\u4e1a"
    return (
        f"{short_name}{report.metadata.report_year}\u5e74\u5ea6\u5916\u8d2d\u7535\u529b\u6392\u653e\u56e0\u5b50"
        "\u91c7\u7528\u751f\u6001\u73af\u5883\u90e8\u53d1\u5e03\u7684\u300a\u5173\u4e8e\u53d1\u5e032022\u5e74"
        "\u7535\u529b\u4e8c\u6c27\u5316\u78b3\u6392\u653e\u56e0\u5b50\u7684\u516c\u544a\u300b\u4e2d\u8981\u6c42\u7684"
        f"2022\u5e74\u5168\u56fd\u7535\u529b\u5e73\u5747\u4e8c\u6c27\u5316\u78b3\u6392\u653e\u56e0\u5b50"
        f"{factors.electricity_factor_tco2_per_mwh} tCO2/MWh\u3002"
    )


def _chapter4_heat_factor_text(report: ReportInput, factors: FactorLibrary) -> str:
    if not report.heat_activity:
        return ""
    return (
        "\u6839\u636e\u300a\u8fd0\u8f93\u4f01\u4e1a\u6838\u7b97\u6307\u5357\u300b\u7684\u89c4\u5b9a\uff0c"
        "\u70ed\u529b\u6392\u653e\u56e0\u5b50\u6ca1\u6709\u5b9e\u6d4b\u503c\uff0c"
        f"\u91c7\u7528\u7f3a\u7701\u503c{factors.heat_factor_tco2_per_gj} tCO2/GJ\u3002"
    )


def _render_calc_fuel_detail(
    document: Document,
    report: ReportInput,
    calculations: CalculationResult,
    factors: FactorLibrary,
) -> None:
    if _skip_empty_table("calc_fuel_detail", bool(calculations.activity_fuel_totals)):
        return
    _add_table_caption(
        document,
        f"\u88685.1-1  {report.metadata.report_year}\u5e74\u71c3\u6599\u71c3\u70e7CO2\u6392\u653e\u91cf\u8ba1\u7b97",
    )
    table = document.add_table(rows=3, cols=6)
    for i, text in enumerate(
        ["\u5e74\u4efd", "\u71c3\u6599\u79cd\u7c7b", "\u6d88\u8017\u91cf", "\u542b\u78b3\u91cf", "\u78b3\u6c27\u5316\u7387", "\u6392\u653e\u91cf"]
    ):
        table.rows[0].cells[i].text = text
    for i, text in enumerate(["", "", "t\u6216\u4e07Nm3", "tC/t\u6216tC/\u4e07Nm3", "-", "tCO2"]):
        table.rows[1].cells[i].text = text
    for i, text in enumerate(["", "", "A", "B", "C", "F=A*B*C*44/12"]):
        table.rows[2].cells[i].text = text
    year = f"{report.metadata.report_year}\u5e74"
    for fuel_type in ("natural_gas", "gasoline", "diesel"):
        if fuel_type not in calculations.activity_fuel_totals:
            continue
        quantity = calculations.activity_fuel_totals[fuel_type]
        factor = factors.fuel_factors[fuel_type]
        cells = table.add_row().cells
        cells[0].text = year
        cells[1].text = _fuel_label(fuel_type)
        cells[2].text = f"{quantity}"
        if factor.emission_factor is not None and (factor.carbon_content == 0 or factor.oxidation_rate == 0):
            cells[3].text = f"{factor.emission_factor}"
            cells[4].text = "—"
        else:
            cells[3].text = str(factor.carbon_content)
            cells[4].text = _format_percent(factor.oxidation_rate)
        cells[5].text = f"{calculations.fuel_emissions.get(fuel_type, 0.0):.2f}"
    total = table.add_row().cells
    total[0].text = "\u5408\u8ba1"
    total[5].text = f"{calculations.fuel_combustion_tco2:.2f}"
    _style_calc_fuel_detail_template_table(table)
    _set_column_vertical_merge(table, 0, 0, 2)
    _set_column_vertical_merge(table, 1, 0, 2)
    if len(table.rows) > 4:
        _set_column_vertical_merge(table, 0, 3, len(table.rows) - 2)
    _merge_row_cells(
        table,
        row_index=len(table.rows) - 1,
        start_col=0,
        end_col=4,
        text="\u5408\u8ba1",
        merged_width="3630",
        bold=False,
    )


def _render_calc_power_heat_detail(
    document: Document,
    report: ReportInput,
    calculations: CalculationResult,
    factors: FactorLibrary,
) -> None:
    _add_table_caption(
        document,
        f"\u88685.2-1  {report.metadata.report_year}\u5e74\u51c0\u8d2d\u5165\u7535\u529b\u548c\u70ed\u529bCO2\u6392\u653e\u91cf\u8ba1\u7b97",
    )
    table = document.add_table(rows=3, cols=5)
    for i, text in enumerate(["\u5e74\u5ea6", "\u9879\u76ee", "\u51c0\u8d2d\u5165\u7535\u529b/\u70ed\u529b", "\u6392\u653e\u56e0\u5b50", "\u6392\u653e\u91cf"]):
        table.rows[0].cells[i].text = text
    for i, text in enumerate(["", "", "MWh\u6216GJ", "tCO2/MWh\u6216tCO2/GJ", "tCO2"]):
        table.rows[1].cells[i].text = text
    for i, text in enumerate(["", "", "A", "B", "C=A*B"]):
        table.rows[2].cells[i].text = text
    year = f"{report.metadata.report_year}\u5e74"
    row = table.add_row().cells
    row[0].text = year
    row[1].text = "\u51c0\u8d2d\u5165\u7535\u529b"
    row[2].text = f"{calculations.electricity_total_mwh}"
    row[3].text = str(factors.electricity_factor_tco2_per_mwh)
    row[4].text = f"{calculations.electricity_emissions_tco2:.2f}"
    row = table.add_row().cells
    row[0].text = year
    row[1].text = "\u51c0\u8d2d\u5165\u70ed\u529b"
    row[2].text = f"{calculations.heat_total_gj}"
    row[3].text = str(factors.heat_factor_tco2_per_gj)
    row[4].text = f"{calculations.heat_emissions_tco2:.2f}"
    total = table.add_row().cells
    total[0].text = "\u5408\u8ba1"
    total[4].text = f"{calculations.electricity_emissions_tco2 + calculations.heat_emissions_tco2:.2f}"
    _style_calc_power_heat_detail_template_table(table)
    _set_column_vertical_merge(table, 0, 0, 2)
    _set_column_vertical_merge(table, 1, 0, 2)
    _set_column_vertical_merge(table, 0, 3, 4)
    _merge_row_cells(
        table,
        row_index=len(table.rows) - 1,
        start_col=0,
        end_col=3,
        text="\u5408\u8ba1",
        merged_width="4084",
        bold=False,
    )


def _render_formula_table(document: Document, role: str) -> None:
    tables: dict[str, list[tuple[str, ...]]] = {
        "formula_total": [
            ("EGHG", "\uff1a", "\u4f01\u4e1a\u6e29\u5ba4\u6c14\u4f53\u78b3\u6392\u653e\u603b\u91cf\uff0ctCO2e\uff1b"),
            ("E\u71c3\u70e7", "\uff1a", "\u51c0\u6d88\u8017\u5404\u79cd\u5316\u77f3\u71c3\u6599\u71c3\u70e7\u6d3b\u52a8\u4ea7\u751f\u7684CO2\u6392\u653e\u91cf\uff0ctCO2\uff1b"),
            ("E\u8fc7\u7a0b", "\uff1a", "\u8fd0\u8f93\u8f66\u8f86\u5728\u5c3e\u6c14\u51c0\u5316\u8fc7\u7a0b\u7531\u4e8e\u4f7f\u7528\u5c3f\u7d20\u7b49\u8fd8\u539f\u5242\u4ea7\u751f\u7684CO2\u6392\u653e\u91cf\uff0ctCO2\uff1b"),
            ("ECO2_\u51c0\u7535", "\uff1a", "\u51c0\u8d2d\u5165\u7535\u529b\u9690\u542b\u7684CO2\u6392\u653e\u91cf\uff0ctCO2\uff1b"),
            ("ECO2_\u51c0\u70ed", "\uff1a", "\u51c0\u8d2d\u5165\u70ed\u529b\u9690\u542b\u7684CO2\u6392\u653e\u91cf\uff0ctCO2\u3002"),
        ],
        "formula_fuel": [
            ("ECO2_\u71c3\u70e7", "\uff1a", "\u5316\u77f3\u71c3\u6599\u71c3\u70e7\u7684CO2\u6392\u653e\u91cf\uff0ctCO2\uff1b"),
            ("i", "\uff1a", "\u5316\u77f3\u71c3\u6599\u7684\u79cd\u7c7b\uff1b"),
            ("ADi,j", "\uff1a", "\u6838\u7b97\u548c\u62a5\u544a\u671f\u5185\u7b2ci\u79cd\u5316\u77f3\u71c3\u6599\u7684\u6d3b\u52a8\u6c34\u5e73\uff0cGJ\uff1b"),
            ("CCi,j", "\uff1a", "\u7b2ci\u79cd\u5316\u77f3\u71c3\u6599\u7684\u5355\u4f4d\u70ed\u503c\u542b\u78b3\u91cf\uff0ctC/GJ\uff1b"),
            ("OFi,j", "\uff1a", "\u7b2ci\u79cd\u5316\u77f3\u71c3\u6599\u7684\u78b3\u6c27\u5316\u7387\uff0c%\uff1b"),
            ("", "\uff1a", "CO2\u4e0eC\u7684\u5206\u5b50\u91cf\u8f6c\u6362\u7cfb\u6570\u3002"),
        ],
        "formula_electricity": [
            ("ECO2_\u51c0\u7535", "\uff1a", "\u4f01\u4e1a\u51c0\u8d2d\u5165\u7684\u7535\u529b\u9690\u542b\u7684CO2\u6392\u653e\uff0ctCO2\uff1b"),
            ("AD\u51c0\u7535", "\uff1a", "\u4f01\u4e1a\u51c0\u8d2d\u5165\u7684\u7535\u529b\u6d88\u8d39\u91cf\uff0cMWh\uff1b"),
            ("EF\u7535\u529b", "\uff1a", "\u5168\u56fd\u7535\u7f51\u5e73\u5747CO2\u6392\u653e\u56e0\u5b50\uff0ctCO2/MWh\uff1b"),
        ],
        "formula_heat": [
            ("ECO2_\u51c0\u70ed", "\uff1a", "\u4f01\u4e1a\u51c0\u8d2d\u5165\u7684\u70ed\u529b\u9690\u542b\u7684CO2\u6392\u653e\uff0ctCO2\uff1b"),
            ("AD\u51c0\u70ed", "\uff1a", "\u4f01\u4e1a\u51c0\u8d2d\u5165\u7684\u70ed\u529b\u6d88\u8d39\u91cf\uff0cGJ\uff1b"),
            ("EF\u70ed\u529b", "\uff1a", "\u70ed\u529b\u4f9b\u5e94\u7684CO2\u6392\u653e\u56e0\u5b50\uff0ctCO2/GJ\u3002"),
        ],
    }
    rows = tables.get(role)
    if not rows:
        return
    table = document.add_table(rows=0, cols=len(rows[0]))
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
    _style_formula_template_table(table, role)
    for row in table.rows:
        _set_cell_omml_formula(row.cells[0], row.cells[0].text.strip())


def _format_quantity(value: float, *, digits: int = 2) -> str:
    text = f"{value:.{digits}f}"
    return text.rstrip("0").rstrip(".") if "." in text else text


def _render_appendix_emission_summary(
    document: Document,
    calculations: CalculationResult,
) -> None:
    table = document.add_table(rows=7, cols=3)
    year = "2024\u5e74"
    rows = [
        [year, "", ""],
        ["\u6e90\u7c7b\u522b", "", "\u6392\u653e\u91cf\uff08tCO2\uff09"],
        [
            "\u5316\u77f3\u71c3\u6599\u71c3\u70e7\u6392\u653e",
            "",
            f"{calculations.fuel_combustion_tco2:.2f}",
        ],
        [
            "\u51c0\u8d2d\u5165\u7535\u529b\u9690\u542b\u6392\u653e\u91cf",
            "",
            f"{calculations.electricity_emissions_tco2:.2f}",
        ],
        [
            "\u51c0\u8d2d\u5165\u70ed\u529b\u9690\u542b\u6392\u653e\u91cf",
            "",
            f"{calculations.heat_emissions_tco2:.2f}",
        ],
        [
            "\u4f01\u4e1a\u6e29\u5ba4\u6c14\u4f53\u6392\u653e\u603b\u91cf",
            "\u5316\u77f3\u71c3\u6599\u71c3\u70e7\u4ea7\u751f\u7684CO2\u6392\u653e",
            f"{calculations.fuel_combustion_tco2:.0f}",
        ],
        [
            "",
            "\u5316\u77f3\u71c3\u6599\u71c3\u70e7\u3001\u51c0\u8d2d\u5165\u7535\u529b\u548c"
            "\u51c0\u8d2d\u5165\u70ed\u529b\u4ea7\u751f\u7684CO2\u6392\u653e",
            f"{calculations.total_emissions_tco2:.0f}",
        ],
    ]
    for row_index, row_values in enumerate(rows):
        for col_index, value in enumerate(row_values):
            table.rows[row_index].cells[col_index].text = value
    _style_appendix_emission_summary_template_table(table)
    _merge_row_cells(
        table,
        row_index=0,
        start_col=0,
        end_col=2,
        text=year,
        merged_width="5000",
    )
    for row_index in (1, 2, 3, 4):
        _merge_row_cells(
            table,
            row_index=row_index,
            start_col=0,
            end_col=1,
            text=rows[row_index][0],
            merged_width="3750",
            bold=row_index < 2,
        )
        _set_cell_width(table.rows[row_index].cells[-1], "2082", "dxa")
    _set_column_vertical_merge(table, 0, 5, 6)
    for row_index in (5, 6):
        _set_cell_width(table.rows[row_index].cells[0], "1446", "pct")
        _set_cell_width(table.rows[row_index].cells[1], "2303", "pct")
        _set_cell_width(table.rows[row_index].cells[2], "2082", "dxa")


def _render_appendix_fuel(
    document: Document,
    report: ReportInput,
    calculations: CalculationResult,
    factors: FactorLibrary,
) -> None:
    if _skip_empty_table("appendix_fuel", bool(calculations.activity_fuel_totals)):
        return
    table = document.add_table(rows=11, cols=6)
    header_rows = [
        [
            "\u71c3\u6599\u54c1\u79cd",
            "\u6d3b\u52a8\u6c34\u5e73",
            "",
            "\u6392\u653e\u56e0\u5b50",
            "",
            "\u6392\u653e\u91cf",
        ],
        [
            "",
            "\u51c0\u6d88\u8017\u91cf",
            "\u4f4e\u4f4d\u53d1\u70ed\u91cf",
            "\u5355\u4f4d\u70ed\u503c\u542b\u78b3\u91cf",
            "\u71c3\u6599\u78b3\u6c27\u5316\u7387",
            "",
        ],
        [
            "",
            "\uff08t\u6216\u4e07Nm3)",
            "(GJ/t,GJ/\u4e07Nm3)",
            "\uff08tC/GJ\uff09",
            "(%)",
            "\uff08tCO2\uff09",
        ],
    ]
    for row_index, row_values in enumerate(header_rows):
        for col_index, value in enumerate(row_values):
            table.rows[row_index].cells[col_index].text = value

    fuel_groups = [
        ("natural_gas", 3, 4),
        ("gasoline", 5, 6),
        ("diesel", 7, 9),
    ]
    for fuel_type, start_row, end_row in fuel_groups:
        quantity = calculations.activity_fuel_totals.get(fuel_type, 0.0)
        factor = factors.fuel_factors.get(fuel_type)
        if factor is None:
            continue
        cells = table.rows[start_row].cells
        cells[0].text = _fuel_label(fuel_type)
        digits = 4 if fuel_type == "natural_gas" else 2
        cells[1].text = _format_quantity(quantity, digits=digits)
        cells[2].text = str(factor.ncv)
        cells[3].text = f"{factor.carbon_content / factor.ncv:.4f}" if factor.ncv else ""
        cells[4].text = f"{factor.oxidation_rate * 100:.0f}"
        cells[5].text = f"{calculations.fuel_emissions.get(fuel_type, 0.0):.2f}"
        for row_index in range(start_row + 1, end_row + 1):
            for col_index in range(6):
                table.rows[row_index].cells[col_index].text = ""
    total = table.rows[10].cells
    total[0].text = "\u5316\u77f3\u71c3\u6599\u71c3\u70e7\u4ea7\u751f\u7684CO2\u6392\u653e\u91cf"
    total[5].text = f"{calculations.fuel_combustion_tco2:.2f}"
    _style_appendix_fuel_template_table(table)
    _set_column_vertical_merge(table, 0, 0, 2)
    _set_column_vertical_merge(table, 5, 0, 2)
    _merge_row_cells(
        table,
        row_index=0,
        start_col=1,
        end_col=2,
        text="\u6d3b\u52a8\u6c34\u5e73",
        merged_width="1691",
    )
    _merge_row_cells(
        table,
        row_index=0,
        start_col=3,
        end_col=4,
        text="\u6392\u653e\u56e0\u5b50",
        merged_width="1642",
    )
    for start_row, end_row in ((3, 4), (5, 6), (7, 9)):
        for col_index in range(6):
            _set_column_vertical_merge(table, col_index, start_row, end_row)
    _merge_row_cells(
        table,
        row_index=10,
        start_col=0,
        end_col=4,
        text="\u5316\u77f3\u71c3\u6599\u71c3\u70e7\u4ea7\u751f\u7684CO2\u6392\u653e\u91cf",
        merged_width="4218",
        bold=False,
    )
    _set_cell_width(table.rows[10].cells[-1], "781", "pct")


def _render_appendix_electricity(
    document: Document,
    calculations: CalculationResult,
    factors: FactorLibrary,
) -> None:
    table = document.add_table(rows=4, cols=4)
    rows = [
        [
            "\u51c0\u8d2d\u5165\u7535\u91cf(MWh)",
            "",
            "\u6392\u653e\u56e0\u5b50\uff08tCO2/MWh\uff09",
            "\u6392\u653e\u91cf(tCO2)",
        ],
        [
            "\u7535\u529b",
            f"{calculations.electricity_total_mwh:.3f}",
            str(factors.electricity_factor_tco2_per_mwh),
            f"{calculations.electricity_emissions_tco2:.2f}",
        ],
        ["\u70ed\u529b", "/", "/", "/"],
        [
            "\u51c0\u8d2d\u5165\u7535\u529b\u9690\u542b\u7684\u4e8c\u6c27\u5316\u78b3\u6392\u653e\u91cf\uff08tCO2\uff09",
            "",
            "",
            f"{calculations.electricity_emissions_tco2:.2f}",
        ],
    ]
    for row_index, row_values in enumerate(rows):
        for col_index, value in enumerate(row_values):
            table.rows[row_index].cells[col_index].text = value
    _style_appendix_electricity_template_table(table)
    _merge_row_cells(
        table,
        row_index=0,
        start_col=0,
        end_col=1,
        text="\u51c0\u8d2d\u5165\u7535\u91cf(MWh)",
        merged_width="2426",
    )
    _merge_row_cells(
        table,
        row_index=3,
        start_col=0,
        end_col=2,
        text="\u51c0\u8d2d\u5165\u7535\u529b\u9690\u542b\u7684\u4e8c\u6c27\u5316\u78b3\u6392\u653e\u91cf\uff08tCO2\uff09",
        merged_width="3933",
        bold=False,
    )
    _set_cell_width(table.rows[3].cells[-1], "1066", "pct")


def _render_appendix_heat(
    document: Document,
    calculations: CalculationResult,
    factors: FactorLibrary,
) -> None:
    table = document.add_table(rows=3, cols=2)
    rows = [
        ["\u51c0\u8d2d\u5165\u70ed\u91cf\uff08GJ\uff09", f"{calculations.heat_total_gj:.2f}"],
        ["\u6392\u653e\u56e0\u5b50\uff08tCO2/GJ\uff09", str(factors.heat_factor_tco2_per_gj)],
        [
            "\u51c0\u8d2d\u5165\u70ed\u529b\u9690\u542b\u7684\u4e8c\u6c27\u5316\u78b3\u6392\u653e\u91cf\uff08tCO2\uff09",
            f"{calculations.heat_emissions_tco2:.2f}",
        ],
    ]
    for row_index, row_values in enumerate(rows):
        for col_index, value in enumerate(row_values):
            table.rows[row_index].cells[col_index].text = value
    _style_appendix_heat_template_table(table)


def _render_table(
    document: Document,
    role: str | None,
    report: ReportInput,
    calculations: CalculationResult,
    factors: FactorLibrary,
) -> None:
    if role in _CUSTOM_TABLE_ROLES and _render_report_table_payload(document, report, role):
        return
    if role == "approval":
        _render_approval(document, report)
    elif role == "equipment":
        _render_equipment(document, report)
    elif role == "workload_monthly":
        _render_workload_fallback(document, report)
    elif role == "emission_sources":
        _render_emission_sources(document, report)
    elif role == "emission_summary":
        _render_emission_summary(document, report, calculations)
    elif role == "appendix_emission_summary":
        _render_appendix_emission_summary(document, calculations)
    elif role == "monitoring":
        _render_monitoring(document, report)
    elif role == "activity_facility":
        _render_activity_facility(document, report)
    elif role == "activity_natural_gas":
        _render_activity_natural_gas(document, report)
    elif role == "activity_fuel":
        _render_activity_fuel(document, report, calculations)
    elif role == "activity_electricity":
        if _has_report_table_role(report, "electricity_internal_detail") or _has_report_table_role(
            report,
            "electricity_balance_detail",
        ):
            return
        _render_activity_electricity_monthly(document, report)
    elif role == "activity_heat":
        _render_activity_heat_sites(document, report)
    elif role == "emission_factors":
        _render_emission_factors(document, report, factors)
    elif role == "calc_fuel_detail":
        _render_calc_fuel_detail(document, report, calculations, factors)
    elif role == "calc_power_heat_detail":
        _render_calc_power_heat_detail(document, report, calculations, factors)
    elif role in {
        "formula_total",
        "formula_fuel",
        "formula_electricity",
        "formula_heat",
    }:
        _render_formula_table(document, role)
    elif role == "appendix_fuel":
        _render_appendix_fuel(document, report, calculations, factors)
    elif role == "appendix_electricity":
        _render_appendix_electricity(document, calculations, factors)
    elif role == "appendix_heat":
        _render_appendix_heat(document, calculations, factors)


def _style_token_for_fixed_key(key: str) -> str:
    if key in {"cover_line1", "cover_line2"}:
        return "cover_title"
    if key == "toc_title":
        return "toc_title"
    if key == "preface.opening_guide":
        return "front_body"
    if key == "preface_transition":
        return "front_transition"
    if key.startswith("method_") and key.endswith("_formula"):
        return "formula_para"
    return "body_para"


def publish_docx(
    report: ReportInput,
    calculations: CalculationResult,
    narratives: dict[str, str],
    output_path: str | Path,
    template: TemplatePackage | None = None,
    chart_png: bytes | None = None,
    blueprint: ReportBlueprint | None = None,
    factors: FactorLibrary | None = None,
) -> Path:
    del template
    factors = factors or FactorLibrary.default_2024()
    library = load_boilerplate_library()
    slots = build_slots(report, factors, calculations)
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    bp = blueprint or load_blueprint()
    toc_entries = build_toc_entries(bp)
    bookmark_by_text_level = {
        (entry.level, entry.text): entry for entry in toc_entries
    }
    # Sequential ids for bookmarkStart/End (must be unique ints in the document)
    next_bookmark_id = 1
    document = Document()
    apply_section_page(document)
    _configure_body_header_style(document)
    fuel_factor_sections_rendered = False
    for block in bp.blocks:
        if block.kind == "cover_chrome":
            _render_cover_chrome(document, report)
        elif block.kind == "heading":
            level = block.level or 1
            text = block.text or ""
            if _is_static_fuel_factor_heading(level, text):
                if not fuel_factor_sections_rendered:
                    next_bookmark_id = _render_fuel_factor_sections(
                        document,
                        report,
                        factors,
                        bookmark_by_text_level,
                        next_bookmark_id,
                    )
                    fuel_factor_sections_rendered = True
                continue
            paragraph = document.add_heading(text, level=level)
            apply_heading_style(paragraph, level)
            entry = bookmark_by_text_level.get((level, text.strip()))
            if entry is not None:
                add_heading_bookmark(paragraph, entry.bookmark, next_bookmark_id)
                next_bookmark_id += 1
            if level == 3 and text.strip().startswith("4.2.1 "):
                grid_factor_text = _chapter4_grid_factor_text(report, factors)
                if grid_factor_text:
                    _add_styled_paragraph(document, grid_factor_text, "body_para")
            elif level == 3 and text.strip().startswith("4.2.2 "):
                heat_factor_text = _chapter4_heat_factor_text(report, factors)
                if heat_factor_text:
                    _add_styled_paragraph(document, heat_factor_text, "body_para")
        elif block.kind == "paragraph_key":
            if block.key == "chapter4_factors":
                continue
            text = narratives.get(block.key or "", "")
            if text.strip():
                _add_styled_paragraphs(document, text, "body_para")
        elif block.kind == "fixed_text":
            text = ""
            key = block.key or ""
            entry = library.by_id.get(key) if key else None
            if entry is None and key:
                append_audit_event("missing_boilerplate_id", {"id": key})
            if entry is not None and requirements_met(entry, report):
                text = render_entry(entry, slots)
            else:
                text = block.text or bp.wording_snippets.get(key, "")
            if text.strip():
                if key.startswith("method_") and key.endswith("_formula"):
                    _add_omml_formula_paragraph(document, key)
                else:
                    _add_styled_paragraph(
                        document, text, _style_token_for_fixed_key(key)
                    )
                    if key == "cover_line1":
                        _add_styled_empty_paragraphs(
                            document,
                            cover_spacer_count("after_title_line1"),
                            "cover_title",
                        )
                    elif key == "cover_line2":
                        add_spacer_paragraphs(
                            document, cover_spacer_count("after_titles")
                        )
                    elif key == "preface_transition":
                        _add_front_blank_paragraphs(
                            document, cover_spacer_count("after_transition")
                        )
                    elif key == "toc_title":
                        _add_front_blank_paragraphs(
                            document, cover_spacer_count("after_toc_title")
                        )
        elif block.kind == "table_role":
            _render_table(document, block.role, report, calculations, factors)
        elif block.kind == "image_slot":
            _render_image(document, report, block.role, block.caption)
        elif (
            block.kind == "chart_slot"
            and block.role == "emissions_distribution"
            and chart_png
        ):
            _add_picture(
                document,
                chart_png,
                width_inches=image_slot_width_inches("emissions_distribution"),
                caption=block.caption,
                align=image_slot_align("emissions_distribution"),
            )
        elif block.kind == "cover_meta":
            _render_cover_meta(document, report)
        elif block.kind == "toc_field":
            try:
                insert_clickable_toc(document, toc_entries)
            except Exception as exc:
                append_audit_event(
                    "toc_field_insert_failed",
                    {"severity": "warning", "error": str(exc)},
                )
        elif block.kind == "page_break":
            document.add_page_break()
        elif block.kind == "section_break":
            _add_section_break(
                document,
                page_number_format=block.page_number_format,
                page_number_start=block.page_number_start,
                section_orientation=block.section_orientation,
                header_text=(
                    report.metadata.header_text
                    if block.page_number_format == "decimal"
                    else None
                ),
            )
    document.save(str(path))
    return path
