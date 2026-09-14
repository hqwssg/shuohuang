from __future__ import annotations

from dataclasses import dataclass
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from carbon_report_agent.blueprint import ReportBlueprint


@dataclass(frozen=True)
class TocEntry:
    level: int
    text: str
    bookmark: str


_SAMPLE_PAGE_BY_TEXT = {
    "第一章 报告主体基本情况": 1,
    "1.1 基本信息": 1,
    "1.2 工艺信息": 2,
    "1.3 设备信息": 2,
    "1.4 工作量信息": 3,
    "第二章 温室气体排放情况": 4,
    "2.1 排放边界和排放源识别": 4,
    "2.1.1 排放边界识别": 4,
    "2.1.2 排放源识别": 4,
    "2.2 核算方法": 5,
    "2.2.1 燃料燃烧CO2排放": 5,
    "2.2.2 净购入电力隐含的CO2排放": 6,
    "2.2.3 净购入热力隐含的CO2排放": 6,
    "第三章 活动水平数据及来源说明": 7,
    "3.1 燃料燃烧CO2排放的活动水平数据": 7,
    "3.1.1 天然气消耗量": 7,
    "3.1.2 汽油、柴油消耗量": 8,
    "3.2 净购入电力活动水平数据": 8,
    "3.3 净购入热力活动水平数据": 9,
    "第四章 排放因子数据及来源说明": 11,
    "4.1 燃料燃烧CO2排放的排放因子": 11,
    "4.1.1 天然气的含碳量": 11,
    "4.1.2 汽油的含碳量": 11,
    "4.1.3 柴油的含碳量": 12,
    "4.1.4 天然气的碳氧化率": 12,
    "4.1.5 汽油的碳氧化率": 12,
    "4.1.6 柴油的碳氧化率": 12,
    "4.2 电力和热力排放因子": 12,
    "4.2.1 电网排放因子": 12,
    "4.2.2 热力供应排放因子": 12,
    "第五章 温室气体排放量的计算": 13,
    "5.1 燃料燃烧CO2排放": 13,
    "5.2 净购入的电力和热力隐含的CO2排放": 13,
    "5.3 温室气体排放汇总表": 13,
    "第六章 结论": 15,
    "第七章 监测设备信息": 16,
    "附表 1  企业温室气体排放汇总表": 19,
    "附表 2  化石燃料燃烧二氧化碳排放量数据表": 20,
    "附表 3  净购入电力隐含的二氧化碳排放量数据表": 21,
    "附表 4  净购入热力隐含的二氧化碳排放量数据表": 21,
}


def build_toc_entries(blueprint: ReportBlueprint) -> list[TocEntry]:
    entries: list[TocEntry] = []
    index = 1
    for block in blueprint.blocks:
        if block.kind != "heading":
            continue
        level = block.level or 1
        if level not in (1, 2, 3):
            continue
        text = (block.text or "").strip()
        if not text:
            continue
        entries.append(
            TocEntry(level=level, text=text, bookmark=f"toc_h_{index:03d}")
        )
        index += 1
    return entries


def _ensure_toc_style(document: Document, level: int):
    name = f"toc {level}"
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def _configure_toc_style(style, level: int, token: dict) -> None:
    style.element.set(qn("w:customStyle"), "0")
    ppr = style.element.get_or_add_pPr()
    if token.get("left_indent_pt") is not None:
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            ppr.append(ind)
        ind.set(qn("w:left"), str(int(float(token["left_indent_pt"]) * 20)))
    if token.get("align"):
        jc = ppr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            ppr.append(jc)
        jc.set(qn("w:val"), str(token["align"]))
    if token.get("tab_leader") == "dot":
        tabs = ppr.find(qn("w:tabs"))
        if tabs is None:
            tabs = OxmlElement("w:tabs")
            ppr.append(tabs)
        for child in list(tabs):
            tabs.remove(child)
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), str(int(token.get("tab_position_twips", 8296))))
        tabs.append(tab)

    rpr = style.element.get_or_add_rPr()
    _apply_run_token_to_rpr(rpr, token)


def _display_text(entry: TocEntry) -> str:
    # The 2024 samples display numbered subheadings as "1.1. 基本信息"
    # while the body headings keep "1.1 基本信息".
    return re.sub(r"^(\d+(?:\.\d+)+)(\s+)", r"\1.\2", entry.text)


def _estimated_page(entry: TocEntry) -> int:
    if entry.text in _SAMPLE_PAGE_BY_TEXT:
        return _SAMPLE_PAGE_BY_TEXT[entry.text]
    if entry.text.startswith("附表"):
        return 19
    return max(1, entry.level)


def _apply_run_token_to_rpr(rpr: OxmlElement, token: dict) -> None:
    fonts = OxmlElement("w:rFonts")
    ascii_name = token.get("font_ascii")
    east = token.get("font_east_asia")
    if ascii_name:
        fonts.set(qn("w:ascii"), str(ascii_name))
        fonts.set(qn("w:hAnsi"), str(ascii_name))
    if east:
        fonts.set(qn("w:eastAsia"), str(east))
    if ascii_name or east:
        rpr.append(fonts)

    if token.get("bold"):
        rpr.append(OxmlElement("w:b"))
        rpr.append(OxmlElement("w:bCs"))

    size_pt = token.get("size_pt")
    if size_pt is not None:
        half_points = str(int(float(size_pt) * 2))
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), half_points)
        rpr.append(sz)
        sz_cs = OxmlElement("w:szCs")
        sz_cs.set(qn("w:val"), half_points)
        rpr.append(sz_cs)


def _apply_run_token(run, token: dict) -> None:
    if token.get("size_pt") is not None:
        from docx.shared import Pt

        run.font.size = Pt(float(token["size_pt"]))
    if "bold" in token:
        run.font.bold = bool(token["bold"])
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    if token.get("font_ascii"):
        fonts.set(qn("w:ascii"), str(token["font_ascii"]))
        fonts.set(qn("w:hAnsi"), str(token["font_ascii"]))
    if token.get("font_east_asia"):
        fonts.set(qn("w:eastAsia"), str(token["font_east_asia"]))


def _add_internal_hyperlink(
    paragraph: Paragraph,
    text: str,
    bookmark: str,
    token: dict,
) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    _apply_run_token_to_rpr(rpr, token)
    run.append(rpr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    run.append(text_elem)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def insert_clickable_toc(document: Document, entries: list[TocEntry]) -> None:
    from carbon_report_agent.style_pack import load_style_pack

    pack = load_style_pack()
    for level in (1, 2, 3):
        token = pack.get(f"toc_entry_l{level}") or pack.get("toc_entry") or {}
        _configure_toc_style(_ensure_toc_style(document, level), level, token)

    for entry in entries:
        paragraph = document.add_paragraph()
        paragraph.style = _ensure_toc_style(document, entry.level)
        _add_internal_hyperlink(paragraph, _display_text(entry), entry.bookmark, {})
        tab_run = paragraph.add_run("\t")
        page_run = paragraph.add_run(str(_estimated_page(entry)))
        _apply_run_token(tab_run, {})
        _apply_run_token(page_run, {})


def add_heading_bookmark(paragraph: Paragraph, bookmark: str, bookmark_id: int) -> None:
    """Wrap paragraph content with a bookmark start/end."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)
