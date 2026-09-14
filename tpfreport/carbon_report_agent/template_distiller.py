from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from pydantic import BaseModel, Field


class TableSummary(BaseModel):
    index: int
    rows: int
    columns: int
    header: list[str]


class TemplatePackage(BaseModel):
    source_path: str
    section_titles: list[str]
    table_count: int
    tables: list[TableSummary] = Field(default_factory=list)
    table_roles: list[str] = Field(default_factory=list)
    wording_snippets: dict[str, str] = Field(default_factory=dict)


DEFAULT_TEMPLATE_PACKAGE_PATH = (
    Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "template_package_suning_2024.json"
)


def load_template_package(path: str | Path | None = None) -> TemplatePackage:
    package_path = Path(path) if path else DEFAULT_TEMPLATE_PACKAGE_PATH
    return TemplatePackage.model_validate_json(package_path.read_text(encoding="utf-8"))


def _clean(text: str) -> str:
    return " ".join(text.split())


def _looks_like_section(text: str) -> bool:
    if not text or len(text) > 40:
        return False
    tokens = [
        "\u57fa\u672c\u60c5\u51b5",
        "\u6392\u653e\u60c5\u51b5",
        "\u6d3b\u52a8\u6c34\u5e73",
        "\u6392\u653e\u56e0\u5b50",
        "\u6392\u653e\u91cf",
        "\u7ed3\u8bba",
        "\u76d1\u6d4b\u8bbe\u5907",
        "\u9644\u5f55",
    ]
    return any(token in text for token in tokens)


def distill_docx_template(path: str | Path) -> TemplatePackage:
    docx_path = Path(path)
    document = Document(str(docx_path))
    section_titles = []
    for paragraph in document.paragraphs:
        text = _clean(paragraph.text)
        if _looks_like_section(text) and text not in section_titles:
            section_titles.append(text)

    tables = []
    for index, table in enumerate(document.tables, start=1):
        header = [_clean(cell.text) for cell in table.rows[0].cells] if table.rows else []
        tables.append(TableSummary(index=index, rows=len(table.rows), columns=len(table.columns), header=header))

    return TemplatePackage(
        source_path=str(docx_path),
        section_titles=section_titles,
        table_count=len(document.tables),
        tables=tables,
        table_roles=["emission-summary", "monitoring-devices", "appendix-fuel", "appendix-electricity", "appendix-heat"],
        wording_snippets={
            "cover_line1": "\u9646\u4e0a\u4ea4\u901a\u8fd0\u8f93\u4f01\u4e1a",
            "cover_line2": "\u6e29\u5ba4\u6c14\u4f53\u6392\u653e\u62a5\u544a",
        },
    )
